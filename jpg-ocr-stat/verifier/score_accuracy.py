"""Original 22-receipt JPG task: offline exact-field accuracy, without answer output.

For data/jpg-ocr-stat-accuracy-20261020. The prediction reader mirrors the
published dataset's evaluation_only/runtime/score_extended.py semantics.
"""

import argparse
import json
import re
from pathlib import Path

from openpyxl import load_workbook


HEADER = ["filename", "date", "total_amount"]
EXPECTED_RECEIPTS = 22
DEFAULT_ACTUAL = Path("/app/workspace/stat_ocr.xlsx")
DEFAULT_ORACLE = Path(__file__).with_name("stat_oracle.xlsx")
DEFAULT_SPEC = Path(__file__).with_name("receipt_spec.json")


def load_spec(spec=None):
    """Use a verifier-side frozen member list, never the submitted row count.

    The original task has no spec file and retains its 22-receipt contract.
    Exported generalization tasks install their fixed 20-member specification.
    """
    if spec is None and not DEFAULT_SPEC.exists():
        return {"expected_receipts": EXPECTED_RECEIPTS, "filenames": None}
    try:
        value = spec if isinstance(spec, dict) else json.loads(
            Path(spec if spec is not None else DEFAULT_SPEC).read_text(encoding="utf-8"))
        count, names = value["expected_receipts"], value["filenames"]
        if type(count) is not int or count <= 0 or not isinstance(names, list) or len(names) != count:
            raise ValueError
        if any(not isinstance(name, str) or not name or Path(name).name != name
               or "/" in name or "\\" in name or name in (".", "..") for name in names):
            raise ValueError
        if len(set(names)) != count or names != sorted(names):
            raise ValueError
        return {"expected_receipts": count, "filenames": names}
    except Exception:
        raise RuntimeError("Frozen receipt specification unavailable or invalid") from None


def cell_string(value):
    return "" if value is None else str(value)


def read_predictions(actual, expected_names):
    predictions, issues, rows = {}, [], []
    if not actual.is_file():
        return predictions, ["output_missing"], rows
    workbook = None
    try:
        workbook = load_workbook(actual, data_only=True)
        if workbook.sheetnames != ["results"]:
            issues.append("invalid_sheets")
        else:
            rows = list(workbook["results"].values)
            if not rows or list(rows[0]) != HEADER:
                issues.append("invalid_header_or_columns")
            else:
                names = [cell_string(row[0]) for row in rows[1:]]
                if len(names) != len(expected_names):
                    issues.append("wrong_row_count")
                if len(names) != len(set(names)):
                    issues.append("duplicate_filenames")
                if set(names) != set(expected_names):
                    issues.append("missing_or_unexpected_filenames")
                if names != sorted(names):
                    issues.append("wrong_filename_order")
                for row in rows[1:]:
                    filename, date, amount = row
                    if not isinstance(filename, str) or not filename:
                        issues.append("invalid_filename_type")
                    # Null is a permitted extraction fallback, scored against GT.
                    if date is not None and (not isinstance(date, str) or not re.fullmatch(r"\d{4}-\d{2}-\d{2}", date)):
                        issues.append("invalid_date_format")
                    if amount is not None and (not isinstance(amount, str) or not re.fullmatch(r"\d+\.\d{2}", amount)):
                        issues.append("invalid_amount_format_or_type")
                    predictions[cell_string(filename)] = [cell_string(date), cell_string(amount)]
    except Exception:
        issues = ["unreadable_workbook"]
        predictions, rows = {}, []
    finally:
        if workbook is not None:
            workbook.close()
    return predictions, sorted(set(issues)), rows


def summarize(group):
    count = len(group)
    dates = sum(row["date_correct"] for row in group)
    amounts = sum(row["amount_correct"] for row in group)
    joint = sum(row["date_correct"] and row["amount_correct"] for row in group)
    return {"receipts": count, "fields": 2 * count, "date_correct": dates, "amount_correct": amounts,
            "field_correct": dates + amounts, "receipt_joint_correct": joint,
            "field_accuracy": (dates + amounts) / (2 * count) if count else None,
            "date_accuracy": dates / count if count else None,
            "amount_accuracy": amounts / count if count else None,
            "receipt_joint_accuracy": joint / count if count else None}


def read_oracle(path, spec=None):
    """Load frozen expected values in memory; reject broken grading fixtures.

    No exception includes cell contents, so an infrastructure failure does not
    disclose the oracle. The original workbook has 22 nonempty string fields
    in each of the two scored columns.
    """
    specification = load_spec(spec)
    expected_count = specification["expected_receipts"]
    workbook = None
    try:
        workbook = load_workbook(path, data_only=True)
        if workbook.sheetnames != ["results"]:
            raise ValueError
        rows = list(workbook["results"].values)
        if not rows or list(rows[0]) != HEADER or len(rows) != expected_count + 1:
            raise ValueError
        if any(len(row) != 3 or not all(isinstance(value, str) and value for value in row) for row in rows[1:]):
            raise ValueError
        names = [row[0] for row in rows[1:]]
        if len(set(names)) != expected_count:
            raise ValueError
        if specification["filenames"] is not None and sorted(names) != specification["filenames"]:
            raise ValueError
        return [(row[0], row[1], row[2]) for row in rows[1:]]
    except Exception:
        raise RuntimeError("Frozen oracle unavailable or invalid; grading did not complete") from None
    finally:
        if workbook is not None:
            workbook.close()


def score(actual, oracle=DEFAULT_ORACLE, spec=None):
    """Return accuracy and boolean feedback; never export expected/actual values."""
    expected_rows = read_oracle(Path(oracle), spec)
    actual = Path(actual)
    predictions, issues, _ = read_predictions(actual, [row[0] for row in expected_rows])
    observations = []
    for filename, date, amount in expected_rows:
        predicted = predictions.get(filename, ["", ""])
        observations.append({"filename": filename,
                             "prediction_present": filename in predictions,
                             "date_correct": predicted[0] == date,
                             "amount_correct": predicted[1] == amount})
    summary = summarize(observations)
    return {"schema_version": 1,
            "task": "jpg-ocr-stat",
            "reward": summary["field_accuracy"],
            "all_receipts": summary,
            "field_feedback": observations,
            "delivery_status": "valid" if not issues else "invalid",
            "delivery_issues": issues,
            "missing_predictions_keep_fixed_denominator": True,
            "field_scoring": "existing_cell_string_and_exact_date_amount_equality_no_new_normalization",
            "format_issues_reported_separately": True,
            "answer_values_exported": False,
            "model_calls": 0}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--actual", type=Path, default=DEFAULT_ACTUAL)
    parser.add_argument("--oracle", type=Path, default=DEFAULT_ORACLE)
    parser.add_argument("--spec", type=Path, help="Frozen receipt membership; omitted for the original 22-image task")
    parser.add_argument("--report", type=Path)
    args = parser.parse_args()
    result = score(args.actual, args.oracle, args.spec)
    text = json.dumps(result, ensure_ascii=False, indent=2) + "\n"
    if args.report:
        args.report.parent.mkdir(parents=True, exist_ok=True)
        args.report.write_text(text, encoding="utf-8")
    else:
        print(text, end="")


if __name__ == "__main__":
    main()
