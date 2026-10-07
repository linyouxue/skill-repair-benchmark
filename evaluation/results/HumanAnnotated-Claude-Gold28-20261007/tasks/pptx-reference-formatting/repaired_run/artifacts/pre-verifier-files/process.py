from copy import deepcopy
from PIL import ImageFont
from pptx import Presentation
from pptx.util import Pt, Emu, Inches
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_AUTO_SIZE
from lxml import etree

SRC = "/root/Awesome-Agent-Papers.pptx"
DST = "/root/Awesome-Agent-Papers_processed.pptx"

NSMAP = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
}
A = "{%s}" % NSMAP["a"]

prs = Presentation(SRC)
slide_w = prs.slide_width
slide_h = prs.slide_height

font_path = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
pil_font = ImageFont.truetype(font_path, 16)

def text_width_emu(text: str) -> int:
    # Measure in pixels at 72 DPI (points), then convert to EMU.
    # DejaVu Sans is a bit wider than Arial, so this gives a safe upper bound.
    bbox = pil_font.getbbox(text)
    px = bbox[2] - bbox[0]
    # getbbox at pt-size returns ~points; 1 pt = 12700 EMU
    emu = int(px * 12700)
    # Add a small horizontal padding for internal text frame margins (~0.1" each side).
    emu += int(Inches(0.2))
    return emu

def set_run_format(run):
    run.font.name = "Arial"
    run.font.size = Pt(16)
    run.font.bold = False
    run.font.color.rgb = RGBColor(0x98, 0x95, 0x96)

# Collect dangling titles (preserve order; dedupe for reference slide)
titles_in_order = []
seen = set()

for slide in prs.slides:
    for shape in list(slide.shapes):
        if shape.is_placeholder or not shape.has_text_frame:
            continue
        txt = shape.text_frame.text.strip()
        if not txt:
            continue

        # Format every run
        for para in shape.text_frame.paragraphs:
            para.alignment = PP_ALIGN.CENTER
            for run in para.runs:
                set_run_format(run)

        # Disable word wrap and auto-size so our explicit width/height is honored.
        shape.text_frame.word_wrap = False
        shape.text_frame.auto_size = MSO_AUTO_SIZE.NONE

        new_w = text_width_emu(txt)
        # Cap width so it doesn't exceed slide width
        max_w = slide_w - Inches(0.2)
        if new_w > max_w:
            new_w = max_w
        shape.width = new_w
        # Keep a sensible single-line height
        shape.height = Inches(0.45)

        # Position at bottom center
        shape.left = (slide_w - shape.width) // 2
        bottom_margin = Inches(0.2)
        shape.top = slide_h - shape.height - bottom_margin

        if txt not in seen:
            seen.add(txt)
            titles_in_order.append(txt)

# Add Reference slide at the end using the same layout as content slides
content_layout = None
for layout in prs.slide_layouts:
    if layout.name == "Title and Content":
        content_layout = layout
        break
if content_layout is None:
    content_layout = prs.slide_layouts[1]

ref_slide = prs.slides.add_slide(content_layout)

# Set title
ref_slide.shapes.title.text = "Reference"

# Find body placeholder
body_ph = None
for ph in ref_slide.placeholders:
    if ph.placeholder_format.idx != 0:
        body_ph = ph
        break

tf = body_ph.text_frame
tf.clear()

# Build paragraphs with auto-numbered bullets via raw XML
txBody = tf._txBody
# Remove existing <a:p> children
for p in txBody.findall(A + "p"):
    txBody.remove(p)

for title in titles_in_order:
    p = etree.SubElement(txBody, A + "p")
    pPr = etree.SubElement(p, A + "pPr")
    pPr.set("lvl", "0")
    # Auto-numbered bullet: 1., 2., 3., ...
    buAutoNum = etree.SubElement(pPr, A + "buAutoNum")
    buAutoNum.set("type", "arabicPeriod")
    r = etree.SubElement(p, A + "r")
    rPr = etree.SubElement(r, A + "rPr")
    rPr.set("lang", "en-US")
    rPr.set("dirty", "0")
    t = etree.SubElement(r, A + "t")
    t.text = title

prs.save(DST)
print(f"Saved to {DST}")
print(f"Collected {len(titles_in_order)} unique titles:")
for i, t in enumerate(titles_in_order, 1):
    print(f"  {i}. {t}")
