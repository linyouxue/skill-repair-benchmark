#!/usr/bin/env python3

from __future__ import annotations

import os
import re
from dataclasses import dataclass
from datetime import date
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from typing import Iterable, Optional

import pytesseract
from PIL import Image, ImageFilter, ImageOps
from openpyxl import Workbook


IMG_DIR = Path("/app/workspace/dataset/img")
OUT_XLSX = Path("/app/workspace/stat_ocr.xlsx")


AMOUNT_RE = re.compile(
    r"(?<![A-Z0-9])"
    r"(\d{1,3}(?:,\d{3})+(?:\.\d{1,2})?|\d+(?:\.\d{1,2})?)"
)

MONTHS = {
    "JAN": 1,
    "FEB": 2,
    "MAR": 3,
    "APR": 4,
    "MAY": 5,
    "JUN": 6,
    "JUL": 7,
    "AUG": 8,
    "SEP": 9,
    "SEPT": 9,
    "OCT": 10,
    "NOV": 11,
    "DEC": 12,
}

EXCLUSION_KEYWORDS = [
    "SUBTOTAL",
    "SUB TOTAL",
    "TAX",
    "GST",
    "SST",
    "DISCOUNT",
    "CHANGE",
    "CASH TENDERED",
]

KEYWORD_GROUPS = [
    ["GRAND TOTAL"],
    ["TOTAL RM", "TOTAL: RM"],
    ["TOTAL AMOUNT"],
    [
        "TOTAL DUE",
        "AMOUNT DUE",
        "BALANCE DUE",
        "NETT TOTAL",
        "NET TOTAL",
        "TOTAL",
        "AMOUNT",
    ],
]


@dataclass(frozen=True)
class TotalCandidate:
    value: str
    priority: int


def list_images_sorted(img_dir: Path) -> list[Path]:
    exts = {".jpg", ".jpeg", ".png", ".webp", ".tif", ".tiff"}
    paths = [p for p in img_dir.iterdir() if p.is_file() and p.suffix.lower() in exts]
    return sorted(paths, key=lambda p: p.name)


def _scale_if_small(img: Image.Image) -> Image.Image:
    w, h = img.size
    if max(w, h) < 1400:
        return img.resize((w * 2, h * 2), Image.Resampling.LANCZOS)
    return img


def ocr_text_variants(img_path: Path) -> list[str]:
    img0 = Image.open(img_path)
    img0 = _scale_if_small(img0)

    gray = ImageOps.grayscale(img0)
    base = ImageOps.autocontrast(gray)

    variants: list[Image.Image] = [
        base,
        base.filter(ImageFilter.SHARPEN),
        ImageOps.invert(base),
        base.point(lambda p: 255 if p > 185 else 0),
        img0,
    ]

    texts: list[str] = []
    for img in variants:
        for psm in (6, 4):
            try:
                txt = pytesseract.image_to_string(img, config=f"--oem 3 --psm {psm}")
            except Exception:
                continue
            if txt and txt.strip():
                texts.append(txt)

    return texts


def normalize_lines(text: str) -> list[str]:
    lines = []
    for ln in text.splitlines():
        ln = ln.strip()
        if not ln:
            continue
        ln = re.sub(r"\s+", " ", ln)
        lines.append(ln)
    return lines


def _contains_any(haystack_upper: str, needles: Iterable[str]) -> bool:
    return any(n in haystack_upper for n in needles)


def parse_amount_str(raw: str) -> Optional[str]:
    s = raw.strip()
    s = s.replace(",", "")
    s = s.replace("O", "0")
    s = s.replace("o", "0")

    if not re.fullmatch(r"\d+(?:\.\d{1,2})?", s):
        return None

    try:
        dec = Decimal(s)
    except InvalidOperation:
        return None

    if dec < 0 or dec > Decimal("10000000"):
        return None

    dec = dec.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    return f"{dec:.2f}"


def find_amounts_in_line(line: str) -> list[str]:
    u = line.upper()
    u = u.replace("RM", " ")
    u = u.replace("MYR", " ")
    u = u.replace("USD", " ")
    u = u.replace("SGD", " ")
    u = u.replace("$", " ")
    found = []
    for m in AMOUNT_RE.finditer(u):
        amt = parse_amount_str(m.group(1))
        if amt is not None:
            found.append(amt)
    return found


def extract_total_from_lines(lines: list[str]) -> Optional[TotalCandidate]:
    if not lines:
        return None

    upper_lines = [ln.upper() for ln in lines]

    for prio, kws in enumerate(KEYWORD_GROUPS):
        for i in range(len(lines) - 1, -1, -1):
            u = upper_lines[i]
            if _contains_any(u, EXCLUSION_KEYWORDS):
                continue
            if not _contains_any(u, kws):
                continue

            amts = find_amounts_in_line(lines[i])
            if amts:
                return TotalCandidate(value=amts[-1], priority=prio)

            if i + 1 < len(lines):
                amts2 = find_amounts_in_line(lines[i + 1])
                if amts2:
                    return TotalCandidate(value=amts2[-1], priority=prio)

    # fallback: last plausible amount on non-excluded lines (avoid obvious tender/change)
    for i in range(len(lines) - 1, -1, -1):
        u = upper_lines[i]
        if _contains_any(u, EXCLUSION_KEYWORDS) or "CASH" in u:
            continue
        amts = find_amounts_in_line(lines[i])
        if amts:
            return TotalCandidate(value=amts[-1], priority=len(KEYWORD_GROUPS))

    return None


def _valid_year(y: int) -> bool:
    return 1990 <= y <= 2099


def _to_year(y: int) -> int:
    if y < 100:
        return 2000 + y if y <= 79 else 1900 + y
    return y


def parse_date_candidate(day: int, month: int, year: int) -> Optional[date]:
    year = _to_year(year)
    if not _valid_year(year):
        return None
    try:
        return date(year, month, day)
    except ValueError:
        return None


def extract_date_from_lines(lines: list[str]) -> Optional[tuple[int, int, str]]:
    best: tuple[int, int, str] | None = None  # (score, -line_index, iso)

    for idx, line in enumerate(lines):
        u = line.upper()
        score_base = 0
        if "DATE" in u or u.startswith("DT") or "DTE" in u:
            score_base += 5
        if idx < 12:
            score_base += 2
        if re.search(r"\b\d{1,2}:\d{2}\b", line):
            score_base += 1

        l = u.replace("O", "0")

        # yyyy-mm-dd or yyyy/mm/dd
        for m in re.finditer(r"\b(\d{4})[\-/\.](\d{1,2})[\-/\.](\d{1,2})\b", l):
            y, mo, d = map(int, m.groups())
            dt = parse_date_candidate(d, mo, y)
            if dt:
                iso = dt.isoformat()
                cand = (score_base + 4, -idx, iso)
                if best is None or cand > best:
                    best = cand

        # dd-mm-yyyy, dd/mm/yy, etc.
        for m in re.finditer(r"\b(\d{1,2})[\-/\.](\d{1,2})[\-/\.](\d{2,4})\b", l):
            a, b, y = map(int, m.groups())
            y = _to_year(y)
            if not _valid_year(y):
                continue

            # disambiguate mm/dd vs dd/mm
            day, mo = a, b
            if a <= 12 and b > 12:
                mo, day = a, b
            elif a > 12 and b <= 12:
                day, mo = a, b
            else:
                # common receipts: dd/mm/yyyy
                day, mo = a, b

            dt = parse_date_candidate(day, mo, y)
            if dt:
                iso = dt.isoformat()
                cand = (score_base + 3, -idx, iso)
                if best is None or cand > best:
                    best = cand

        # dd MON yyyy
        for m in re.finditer(r"\b(\d{1,2})\s*([A-Z]{3,4})\s*(\d{2,4})\b", l):
            d = int(m.group(1))
            mon = m.group(2)
            y = int(m.group(3))
            if mon not in MONTHS:
                continue
            dt = parse_date_candidate(d, MONTHS[mon], y)
            if dt:
                iso = dt.isoformat()
                cand = (score_base + 2, -idx, iso)
                if best is None or cand > best:
                    best = cand

    return best


def extract_fields(img_path: Path) -> tuple[Optional[str], Optional[str]]:
    texts = ocr_text_variants(img_path)

    best_total: TotalCandidate | None = None
    best_date: tuple[int, int, str] | None = None

    for t in texts:
        lines = normalize_lines(t)
        if not lines:
            continue

        total = extract_total_from_lines(lines)
        if total is not None:
            if best_total is None or total.priority < best_total.priority:
                best_total = total

        d = extract_date_from_lines(lines)
        if d is not None:
            if best_date is None or d > best_date:
                best_date = d

        if best_total is not None and best_total.priority == 0 and best_date is not None:
            break

    return (best_date[2] if best_date else None, best_total.value if best_total else None)


def write_results(rows: list[tuple[str, Optional[str], Optional[str]]], out_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "results"

    ws.cell(row=1, column=1, value="filename")
    ws.cell(row=1, column=2, value="date")
    ws.cell(row=1, column=3, value="total_amount")

    for r, (fn, dt, amt) in enumerate(rows, start=2):
        ws.cell(row=r, column=1, value=fn)
        ws.cell(row=r, column=2, value=dt)
        ws.cell(row=r, column=3, value=amt)

    wb.save(out_path)


def main() -> int:
    images = list_images_sorted(IMG_DIR)
    rows: list[tuple[str, Optional[str], Optional[str]]] = []

    for p in images:
        dt, amt = extract_fields(p)
        rows.append((p.name, dt, amt))

    write_results(rows, OUT_XLSX)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
