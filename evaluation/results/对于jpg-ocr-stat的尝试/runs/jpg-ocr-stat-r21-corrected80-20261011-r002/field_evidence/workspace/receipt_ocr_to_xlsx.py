import re
from dataclasses import dataclass
from datetime import date
from pathlib import Path

import pytesseract
from PIL import Image, ImageOps
from openpyxl import Workbook, load_workbook


IMG_DIR = Path("/app/workspace/dataset/img")
OUT_XLSX = Path("/app/workspace/stat_ocr.xlsx")


@dataclass(frozen=True)
class Line:
    text: str
    left: int
    top: int
    right: int
    bottom: int


def _lines_from_tesseract(img: Image.Image, *, config: str) -> list[Line]:
    data = pytesseract.image_to_data(img, output_type=pytesseract.Output.DICT, config=config)

    groups: dict[tuple[int, int, int], dict] = {}
    n = len(data.get("text", []))
    for i in range(n):
        word = (data["text"][i] or "").strip()
        if not word:
            continue
        key = (int(data["block_num"][i]), int(data["par_num"][i]), int(data["line_num"][i]))
        left = int(data["left"][i])
        top = int(data["top"][i])
        width = int(data["width"][i])
        height = int(data["height"][i])
        right = left + width
        bottom = top + height

        g = groups.get(key)
        if g is None:
            groups[key] = {
                "words": [word],
                "left": left,
                "top": top,
                "right": right,
                "bottom": bottom,
            }
        else:
            g["words"].append(word)
            g["left"] = min(g["left"], left)
            g["top"] = min(g["top"], top)
            g["right"] = max(g["right"], right)
            g["bottom"] = max(g["bottom"], bottom)

    lines: list[Line] = []
    for g in groups.values():
        text = " ".join(g["words"]).strip()
        if text:
            lines.append(Line(text=text, left=g["left"], top=g["top"], right=g["right"], bottom=g["bottom"]))

    lines.sort(key=lambda l: (l.top, l.left))
    return lines


def ocr_variants(image_path: Path) -> list[list[Line]]:
    img = Image.open(image_path)
    gray = ImageOps.grayscale(img)

    variants: list[Image.Image] = []
    variants.append(ImageOps.autocontrast(gray))
    variants.append(ImageOps.invert(ImageOps.autocontrast(gray)))

    up2 = gray.resize((gray.width * 2, gray.height * 2))
    variants.append(ImageOps.autocontrast(up2))
    variants.append(ImageOps.invert(ImageOps.autocontrast(up2)))

    cfg = "--oem 3 --psm 6"
    return [_lines_from_tesseract(v, config=cfg) for v in variants]


_AMOUNT_RE = re.compile(r"(?<!\d)(\d{1,3}(?:,\d{3})*|\d+)\.(\d{2})(?!\d)")


def _normalize_amount(num: str) -> str | None:
    try:
        v = float(num.replace(",", ""))
    except ValueError:
        return None
    if v < 0:
        return None
    return f"{v:.2f}"


_TOTAL_RULES: list[tuple[re.Pattern[str], int]] = [
    (re.compile(r"\bGRAND\s+TOTAL\b", re.I), 100),
    (re.compile(r"\bTOTAL\s*RM\b|\bTOTAL\s*:\s*RM\b", re.I), 90),
    (re.compile(r"\bTOTAL\s+AMOUNT\b", re.I), 80),
    (re.compile(r"\bNETT\s+TOTAL\b|\bNET\s+TOTAL\b", re.I), 75),
    (re.compile(r"\bTOTAL\s+(?:DUE|ROUNDED)\b|\bAMOUNT\s+DUE\b|\bBALANCE\s+DUE\b", re.I), 70),
    (re.compile(r"\bTOTAL\b|\bAMOUNT\b", re.I), 60),
]


def extract_total_amount(lines: list[Line]) -> tuple[str | None, dict]:
    best: tuple[int, float, str] | None = None  # (score, top, amount)
    evidence: dict = {"source_line": None, "matched_rule": None}

    for i, line in enumerate(lines):
        t = line.text
        tu = t.upper()

        rule_score = None
        rule_pat = None
        for pat, score in _TOTAL_RULES:
            if pat.search(t):
                rule_score = score
                rule_pat = pat.pattern
                break
        if rule_score is None:
            continue

        if "SUBTOTAL" in tu or "SUB TOTAL" in tu:
            rule_score -= 80

        amounts = [m.group(0) for m in _AMOUNT_RE.finditer(t)]
        picked = amounts[-1] if amounts else None

        used_next_line = False
        if picked is None and i + 1 < len(lines):
            t2 = lines[i + 1].text.strip()
            t2u = t2.upper()
            if (
                (t.rstrip().endswith(":") or tu.endswith("TOTAL") or tu.endswith("TOTAL:"))
                and not any(p.search(t2) for p, _ in _TOTAL_RULES)
                and not any(k in t2u for k in ("SUBTOTAL", "SUB TOTAL", "TAX", "GST", "SST", "DISCOUNT", "CHANGE", "CASH"))
            ):
                amounts2 = [m.group(0) for m in _AMOUNT_RE.finditer(t2)]
                if amounts2 and len(t2.split()) <= 3:
                    picked = amounts2[-1]
                    rule_score -= 5
                    used_next_line = True

        if picked is None:
            continue

        norm = _normalize_amount(picked)
        if norm is None:
            continue

        bottomness = line.top / max(1, lines[-1].bottom)
        score = rule_score + int(10 * bottomness)
        if used_next_line:
            score -= 1

        candidate = (score, float(line.top), norm)
        if best is None or candidate > best:
            best = candidate
            evidence = {"source_line": t, "matched_rule": rule_pat}

    return (best[2] if best else None), evidence


_MONTHS = {
    "JAN": 1,
    "FEB": 2,
    "MAR": 3,
    "APR": 4,
    "MAY": 5,
    "JUN": 6,
    "JUL": 7,
    "AUG": 8,
    "SEP": 9,
    "SEPT": 9,
    "OCT": 10,
    "NOV": 11,
    "DEC": 12,
}

_DATE_PATTERNS: list[re.Pattern[str]] = [
    re.compile(r"(?<!\d)(\d{4})[\-/](\d{1,2})[\-/](\d{1,2})(?!\d)"),
    re.compile(r"(?<!\d)(\d{1,2})[\-/](\d{1,2})[\-/](\d{2,4})(?!\d)"),
    re.compile(r"(?<!\w)(\d{1,2})\s*(JAN|FEB|MAR|APR|MAY|JUN|JUL|AUG|SEP|SEPT|OCT|NOV|DEC)[A-Z]*\s*(\d{2,4})(?!\w)", re.I),
    re.compile(r"(?<!\w)(JAN|FEB|MAR|APR|MAY|JUN|JUL|AUG|SEP|SEPT|OCT|NOV|DEC)[A-Z]*\s*(\d{1,2})\s*,?\s*(\d{2,4})(?!\w)", re.I),
]


def _to_iso(y: int, m: int, d: int) -> str | None:
    try:
        return date(y, m, d).isoformat()
    except ValueError:
        return None


def extract_date(lines: list[Line]) -> tuple[str | None, dict]:
    best: tuple[int, float, str] | None = None  # (score, top, iso)
    evidence: dict = {"source_line": None, "pattern": None}

    for line in lines:
        t = line.text
        tu = t.upper()
        for pat in _DATE_PATTERNS:
            for m in pat.finditer(t):
                iso = None
                score = 0

                if pat is _DATE_PATTERNS[0]:
                    y, mo, da = int(m.group(1)), int(m.group(2)), int(m.group(3))
                    iso = _to_iso(y, mo, da)
                    score += 8
                elif pat is _DATE_PATTERNS[1]:
                    a, b, c = int(m.group(1)), int(m.group(2)), int(m.group(3))
                    y = c + 2000 if c < 100 else c
                    if a > 12 and b <= 12:
                        d, mo = a, b
                    elif b > 12 and a <= 12:
                        mo, d = a, b
                    else:
                        d, mo = a, b
                    iso = _to_iso(y, mo, d)
                    score += 6
                elif pat is _DATE_PATTERNS[2]:
                    d = int(m.group(1))
                    mo = _MONTHS[m.group(2).upper()]
                    y = int(m.group(3))
                    y = y + 2000 if y < 100 else y
                    iso = _to_iso(y, mo, d)
                    score += 7
                else:
                    mo = _MONTHS[m.group(1).upper()]
                    d = int(m.group(2))
                    y = int(m.group(3))
                    y = y + 2000 if y < 100 else y
                    iso = _to_iso(y, mo, d)
                    score += 7

                if iso is None:
                    continue

                if "DATE" in tu:
                    score += 10
                if "TIME" in tu:
                    score += 1
                if re.search(r"\b\d{2}:\d{2}(?::\d{2})?\b", t):
                    score += 1
                if m.group(0) and len(m.group(0)) >= 10:
                    score += 2

                candidate = (score, float(line.top), iso)
                if best is None or candidate > best:
                    best = candidate
                    evidence = {"source_line": t, "pattern": pat.pattern}

    return (best[2] if best else None), evidence


def best_field_from_variants(variants: list[list[Line]], extractor):
    best_val = None
    best_score = -1
    best_evidence = None

    for lines in variants:
        val, ev = extractor(lines)
        if val is None:
            continue
        score = 1
        if ev.get("source_line"):
            score += 1
        if extractor is extract_total_amount and ev.get("matched_rule"):
            score += 1
        if extractor is extract_date and "DATE" in (ev.get("source_line") or "").upper():
            score += 2

        if score > best_score:
            best_score = score
            best_val = val
            best_evidence = ev

    return best_val, best_evidence


def write_xlsx(rows: list[tuple[str, str | None, str | None]]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "results"

    ws.append(["filename", "date", "total_amount"])
    for fn, d, amt in rows:
        ws.append([fn, d, amt])

    wb.save(OUT_XLSX)


def verify_xlsx(expected_rows: list[tuple[str, str | None, str | None]]) -> None:
    wb = load_workbook(OUT_XLSX)
    assert wb.sheetnames == ["results"], f"unexpected sheets: {wb.sheetnames}"
    ws = wb["results"]

    got = []
    for r in ws.iter_rows(values_only=True):
        got.append(r)

    assert got[0] == ("filename", "date", "total_amount"), f"bad header: {got[0]}"
    got_rows = [(a, b, c) for (a, b, c) in got[1:]]
    assert got_rows == expected_rows, "xlsx content mismatch"


def main() -> None:
    images = sorted([p for p in IMG_DIR.iterdir() if p.is_file()])
    rows: list[tuple[str, str | None, str | None]] = []

    for p in images:
        variants = ocr_variants(p)
        d, _ = best_field_from_variants(variants, extract_date)
        amt, _ = best_field_from_variants(variants, extract_total_amount)
        rows.append((p.name, d, amt))

    write_xlsx(rows)
    verify_xlsx(rows)


if __name__ == "__main__":
    main()
