#!/usr/bin/env python3
from __future__ import annotations

import re
from dataclasses import dataclass
from datetime import date
from pathlib import Path
from typing import Iterable, Optional

import pytesseract
from PIL import Image, ImageFilter, ImageOps
from openpyxl import Workbook


IMG_DIR = Path("/app/workspace/dataset/img")
OUT_XLSX = Path("/app/workspace/stat_ocr.xlsx")


MONTHS = {
    "JAN": 1,
    "FEB": 2,
    "MAR": 3,
    "APR": 4,
    "MAY": 5,
    "JUN": 6,
    "JUL": 7,
    "AUG": 8,
    "SEP": 9,
    "SEPT": 9,
    "OCT": 10,
    "NOV": 11,
    "DEC": 12,
}


def _safe_date(y: int, m: int, d: int) -> Optional[date]:
    try:
        return date(y, m, d)
    except Exception:
        return None


def _normalize_year(y: int) -> int:
    if y < 100:
        return 2000 + y if y < 80 else 1900 + y
    return y


def _preprocess(img: Image.Image, *, upscale: int = 1, unsharp: bool = False) -> Image.Image:
    gray = ImageOps.grayscale(img)
    gray = ImageOps.autocontrast(gray)

    if upscale > 1:
        gray = gray.resize((gray.size[0] * upscale, gray.size[1] * upscale), Image.Resampling.LANCZOS)

    if unsharp:
        gray = gray.filter(ImageFilter.UnsharpMask(radius=2, percent=250, threshold=3))
    else:
        gray = gray.filter(ImageFilter.SHARPEN)

    return gray


def _ocr_best(img: Image.Image, *, aggressive: bool = False) -> str:
    w, _h = img.size
    if aggressive:
        upscale = 3 if w < 700 else 2 if w < 1200 else 1
        processed = _preprocess(img, upscale=upscale, unsharp=True)
    else:
        upscale = 2 if max(img.size) < 1400 else 1
        processed = _preprocess(img, upscale=upscale, unsharp=False)

    best_text = ""
    best_score = -1
    for psm in (4, 6):
        txt = pytesseract.image_to_string(processed, config=f"--psm {psm}")
        score = sum(ch.isalnum() for ch in txt)
        if score > best_score:
            best_score = score
            best_text = txt

    return best_text


def _ocr_date_special(img: Image.Image) -> str:
    w, h = img.size
    crop = img.crop((0, int(h * 0.20), w, int(h * 0.60)))
    gray = ImageOps.autocontrast(ImageOps.grayscale(crop))
    gray = gray.resize((gray.size[0] * 3, gray.size[1] * 3), Image.Resampling.LANCZOS)
    texts = [pytesseract.image_to_string(gray, config=f"--psm {psm}") for psm in (6, 7)]
    return "\n".join(texts)



@dataclass
class DateCandidate:
    value: date
    score: int
    line_index: int


def extract_date(text: str) -> Optional[str]:
    lines = [ln.strip() for ln in text.splitlines()]
    candidates: list[DateCandidate] = []

    iso_re = re.compile(r"\b(20\d{2}|19\d{2})[\-/.](\d{1,2})[\-/.](\d{1,2})\b")
    dmy_re = re.compile(r"\b(\d{1,2})[\-/.](\d{1,2})[\-/.]((?:20)?\d{2}|19\d{2})\b")
    d_mon_re = re.compile(
        r"\b(\d{1,2})\s*(JAN|FEB|MAR|APR|MAY|JUN|JUL|AUG|SEP|SEPT|OCT|NOV|DEC)[A-Z]*\s*(\d{2,4})\b",
        re.I,
    )
    mon_d_re = re.compile(
        r"\b(JAN|FEB|MAR|APR|MAY|JUN|JUL|AUG|SEP|SEPT|OCT|NOV|DEC)[A-Z]*\s*(\d{1,2})\s*,?\s*(\d{2,4})\b",
        re.I,
    )

    for i, raw in enumerate(lines):
        if not raw:
            continue
        upper = raw.upper()

        base = 0
        if any(k in upper for k in ("DATE", "OATE", "OATES", "DATS")):
            base += 50
            if "PRINT" in upper:
                base += 10
            if "PRINT DATE" in upper:
                base -= 25
        if "RECEIPT" in upper:
            base += 25
        if "GST" in upper and not any(k in upper for k in ("DATE", "OATE", "OATES", "DATS")):
            base -= 10

        for y, m, d in iso_re.findall(upper):
            dt = _safe_date(int(y), int(m), int(d))
            if dt:
                candidates.append(DateCandidate(dt, base + 12, i))

        for a, b, y in dmy_re.findall(upper):
            yy = _normalize_year(int(y))
            aa = int(a)
            bb = int(b)
            # Prefer day-first unless evidence suggests month-first.
            if aa > 12:
                d, m = aa, bb
            elif bb > 12:
                m, d = aa, bb
            else:
                d, m = aa, bb
            dt = _safe_date(yy, m, d)
            if dt:
                candidates.append(DateCandidate(dt, base + 8, i))

        for d, mon, y in d_mon_re.findall(upper):
            yy = _normalize_year(int(y))
            mm = MONTHS.get(mon[:4].upper(), MONTHS.get(mon[:3].upper()))
            if mm:
                dt = _safe_date(yy, mm, int(d))
                if dt:
                    candidates.append(DateCandidate(dt, base + 14, i))

        for mon, d, y in mon_d_re.findall(upper):
            yy = _normalize_year(int(y))
            mm = MONTHS.get(mon[:4].upper(), MONTHS.get(mon[:3].upper()))
            if mm:
                dt = _safe_date(yy, mm, int(d))
                if dt:
                    candidates.append(DateCandidate(dt, base + 10, i))

    if not candidates:
        return None

    candidates.sort(key=lambda c: (-c.score, c.line_index))
    return candidates[0].value.isoformat()


def _normalize_label_text(s: str) -> str:
    u = s.upper()
    u = u.replace("0", "O")
    u = re.sub(r"TO[IL]AL", "TOTAL", u)
    u = re.sub(r"T[OA]TAL", "TOTAL", u)
    u = re.sub(r"\bFATAL\b", "TOTAL", u)
    return u


def _money_candidates(s: str, *, allow_implicit_cents: bool = False) -> list[str]:
    s = s.replace("O", "0")
    s = s.replace("o", "0")
    s = re.sub(r"(?<=\d)\s+(?=[\.,])", "", s)
    s = re.sub(r"(?<=[\.,])\s+(?=\d)", "", s)

    out: list[str] = []
    for m in re.finditer(r"(?<!\d)(\d{1,3}(?:,\d{3})*|\d+)[\.,](\d{2})\b", s):
        out.append(f"{m.group(1).replace(',', '')}.{m.group(2)}")

    if out:
        return out

    for m in re.finditer(r"(?<!\d)(\d{1,3}(?:,\d{3})*|\d+)[\.,](\d)\b", s):
        out.append(f"{m.group(1).replace(',', '')}.{m.group(2)}0")

    if out:
        return out

    if allow_implicit_cents:
        ints = re.findall(r"(?<!\d)(\d{3,7})(?!\d)", s)
        for n in ints:
            iv = int(n)
            if 1900 <= iv <= 2099:
                continue
            out.append(f"{iv / 100:.2f}")
        if out:
            return out

    return []


def extract_total_amount(text: str) -> Optional[str]:
    raw_lines = [ln.rstrip() for ln in text.splitlines()]
    lines = [re.sub(r"\s+", " ", ln).strip() for ln in raw_lines]

    keyword_groups: list[tuple[int, tuple[str, ...]]] = [
        (0, ("GRAND TOTAL",)),
        (1, ("TOTAL RM", "TOTAL: RM", "TOTAL : RM")),
        (2, ("TOTAL AMOUNT", "TOTAL SALES")),
        (3, ("TOTAL AFTER ROUNDING",)),
        (4, ("NETT TOTAL", "NET TOTAL")),
        (5, ("TOTAL DUE", "AMOUNT DUE", "BALANCE DUE")),
        (6, ("TOTAL",)),
    ]

    non_payable = (
        "SUBTOTAL",
        "SUB TOTAL",
        "SUB-TOTAL",
        "DISCOUNT",
        "CHANGE",
        "CASH TENDERED",
        "CASH RECEIVED",
        "PAYMENT",
        "ITEM COUNT",
        "TOTAL QUANTITY",
        "QTY",
        "QUANTITY",
    )

    skip_context = ("GST SUMMARY", "TAX SUMMARY", "GST RATE", "GST SUMMARY AMOUNT")

    candidates: list[tuple[int, int, int, str]] = []  # (priority, -score, -line_index, amount)

    for idx, ln in enumerate(lines):
        if not ln:
            continue
        upper = ln.upper()
        label_text = _normalize_label_text(upper)

        if any(k in label_text for k in ("SUBTOTAL", "SUB TOTAL", "SUB-TOTAL", "TOTAL QUANTITY")):
            continue
        if any(k in label_text for k in skip_context):
            continue

        matched_priority: Optional[int] = None
        matched_kw: Optional[str] = None
        for pri, kws in keyword_groups:
            for kw in kws:
                if kw in label_text:
                    matched_priority = pri
                    matched_kw = kw
                    break
            if matched_priority is not None:
                break

        if matched_priority is None:
            continue

        if matched_priority >= 6 and any(bad in label_text for bad in non_payable):
            continue
        if matched_priority >= 6 and "TOTAL AFTER" in label_text:
            matched_priority = 3

        amts = _money_candidates(upper, allow_implicit_cents=True)
        if not amts and idx + 1 < len(lines):
            nxt = lines[idx + 1].strip()
            nxt_u = nxt.upper()
            nxt_amts = _money_candidates(nxt_u, allow_implicit_cents=True)
            if nxt_amts:
                # Only accept split lines if the next line is basically just an amount.
                stripped = re.sub(r"[0-9\s,\.]", "", nxt_u)
                stripped = stripped.replace("RM", "")
                if len(stripped.strip()) <= 2:
                    amts = nxt_amts

        if not amts:
            continue

        amt = amts[-1]
        try:
            val = float(amt)
        except Exception:
            continue
        if not (0 <= val < 1e9):
            continue

        score = 0
        if matched_priority == 0:
            score += 20
        elif matched_priority == 1:
            score += 12
        elif matched_priority == 2:
            score += 10
        elif matched_priority == 3:
            score += 14
        elif matched_priority == 4:
            score += 9
        elif matched_priority == 5:
            score += 11

        if any(k in label_text for k in ("INCL", "INCLUSIVE", "INCLUDED", "INC]")):
            score += 4
        if "EXCL" in label_text:
            score -= 10
        if ("ROUNDED" in label_text or "AMT ROUNDED" in label_text) and matched_priority != 3:
            score -= 20

        candidates.append((matched_priority, -score, -idx, f"{val:.2f}"))

    if not candidates:
        return None

    candidates.sort(key=lambda t: (t[0], t[1], t[2]))
    return candidates[0][3]


def iter_images(img_dir: Path) -> Iterable[Path]:
    exts = {".jpg", ".jpeg", ".png", ".webp"}
    files = [p for p in img_dir.iterdir() if p.is_file() and p.suffix.lower() in exts]
    return sorted(files, key=lambda p: p.name)


def write_xlsx(rows: list[tuple[str, Optional[str], Optional[str]]], out_path: Path) -> None:
    wb = Workbook()
    default = wb.active
    wb.remove(default)

    ws = wb.create_sheet("results")
    ws.append(["filename", "date", "total_amount"])
    for fn, dt, amt in rows:
        ws.append([fn, dt, amt])

    wb.save(out_path)


def verify_xlsx(out_path: Path, expected_rows: int) -> None:
    from openpyxl import load_workbook

    wb = load_workbook(out_path)
    assert wb.sheetnames == ["results"], wb.sheetnames
    ws = wb["results"]
    assert ws.max_row == expected_rows + 1
    assert ws.max_column == 3
    header = [ws.cell(row=1, column=i).value for i in range(1, 4)]
    assert header == ["filename", "date", "total_amount"], header


def main() -> None:
    rows: list[tuple[str, Optional[str], Optional[str]]] = []

    for img_path in iter_images(IMG_DIR):
        with Image.open(img_path) as im:
            w, h = im.size
            text = _ocr_best(im)

            dt_base = extract_date(text)
            header = im.crop((0, int(h * 0.15), w, int(h * 0.45)))
            dt_header = extract_date(_ocr_best(header))
            dt = dt_header or dt_base
            if dt is None:
                dt = extract_date(_ocr_date_special(im))

            total = extract_total_amount(text)
            if total is None:
                bottom = im.crop((0, int(h * 0.40), w, h))
                total = extract_total_amount(_ocr_best(bottom, aggressive=True))

        rows.append((img_path.name, dt, total))

    write_xlsx(rows, OUT_XLSX)
    verify_xlsx(OUT_XLSX, expected_rows=len(rows))


if __name__ == "__main__":
    main()
