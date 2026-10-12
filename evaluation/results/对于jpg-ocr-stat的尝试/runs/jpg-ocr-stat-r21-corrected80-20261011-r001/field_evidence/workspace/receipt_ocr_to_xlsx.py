#!/usr/bin/env python3

from __future__ import annotations

import json
import re
from dataclasses import dataclass
from datetime import date
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from typing import Any, Iterable

from PIL import Image, ImageEnhance, ImageFilter, ImageOps
import pytesseract
from openpyxl import Workbook, load_workbook


IMG_DIR = Path("/app/workspace/dataset/img")
OUT_XLSX = Path("/app/workspace/stat_ocr.xlsx")
WORK_DIR = Path("/app/workspace/ocr_work")
LEDGER_PATH = WORK_DIR / "field_ledger.json"
QUEUE_PATH = WORK_DIR / "review_queue.json"

FIELDS = ("date", "total_amount")

PAYABLE_KEYWORDS_PRIORITY = [
    "GRAND TOTAL",
    "TOTAL RM",
    "TOTAL: RM",
    "TOTAL AMOUNT",
    "TOTAL DUE",
    "AMOUNT DUE",
    "BALANCE DUE",
    "NETT TOTAL",
    "NET TOTAL",
    "TOTAL",
    "AMOUNT",
]

NON_PAYABLE_CUES = [
    "SUBTOTAL",
    "SUB TOTAL",
    "TAX",
    "GST",
    "SST",
    "DISCOUNT",
    "CHANGE",
    "CASH TENDERED",
    "CASH",
]


def enumerate_input_images_sorted() -> list[Path]:
    exts = {".jpg", ".jpeg", ".png", ".webp"}
    imgs = [p for p in IMG_DIR.iterdir() if p.is_file() and p.suffix.lower() in exts]
    return sorted(imgs, key=lambda p: p.name)


def _default_field_record() -> dict[str, Any]:
    return {
        "value": None,
        "state": "unresolved",
        "first_pass_done": False,
        "candidates": [],
        "evidence": {
            "raw_text": None,
            "label": None,
            "bbox": None,
            "ocr_source": None,
        },
        "confirmation_reason": None,
        "review_reason": "first pass pending",
        "changes": [],
    }


def load_working_records_or_initialize(images: list[Path]) -> dict[str, dict[str, Any]]:
    WORK_DIR.mkdir(parents=True, exist_ok=True)

    if LEDGER_PATH.exists():
        with LEDGER_PATH.open("r", encoding="utf-8") as f:
            records = json.load(f)
    else:
        records = {}

    for p in images:
        records.setdefault(p.name, {})
        for field in FIELDS:
            records[p.name].setdefault(field, _default_field_record())

    # Write back to ensure schema completeness
    with LEDGER_PATH.open("w", encoding="utf-8") as f:
        json.dump(records, f, indent=2, ensure_ascii=False)

    # Reopen and print counts
    with LEDGER_PATH.open("r", encoding="utf-8") as f:
        reopened = json.load(f)

    total_fields = sum(len(v) for v in reopened.values())
    states = {}
    for fn in reopened:
        for field in FIELDS:
            st = reopened[fn][field]["state"]
            states[st] = states.get(st, 0) + 1

    print(
        f"ledger loaded: images={len(reopened)} fields={total_fields} states={states}",
        flush=True,
    )
    return reopened


def save_working_records(records: dict[str, Any]) -> None:
    with LEDGER_PATH.open("w", encoding="utf-8") as f:
        json.dump(records, f, indent=2, ensure_ascii=False)


@dataclass(frozen=True)
class OCRPass:
    variant: str
    psm: int
    text: str


def _prep_variants(img: Image.Image) -> list[tuple[str, Image.Image]]:
    base = ImageOps.grayscale(img)
    variants: list[tuple[str, Image.Image]] = []

    ac = ImageOps.autocontrast(base)
    variants.append(("gray_autocontrast", ac))

    inv = ImageOps.invert(ac)
    variants.append(("invert_autocontrast", inv))

    sharp = ac.filter(ImageFilter.SHARPEN)
    variants.append(("sharpen", sharp))

    # Slightly boost contrast
    enh = ImageEnhance.Contrast(ac).enhance(1.6)
    variants.append(("contrast_1.6", enh))

    # Thresholds
    for t in (120, 150):
        thr = ac.point(lambda p, tt=t: 255 if p > tt else 0)
        variants.append((f"threshold_{t}", thr))

    # Upscale
    scale = 2
    up = ac.resize((ac.width * scale, ac.height * scale), Image.Resampling.LANCZOS)
    variants.append(("upscale2x", up))

    return variants


def recognize_with_text_and_regions(image_path: Path, crop_bbox: tuple[int, int, int, int] | None = None) -> list[OCRPass]:
    img = Image.open(image_path)
    if crop_bbox is not None:
        img = img.crop(crop_bbox)

    passes: list[OCRPass] = []
    for variant_name, variant_img in _prep_variants(img):
        for psm in (6, 4, 11):
            config = f"--oem 3 --psm {psm}"
            txt = pytesseract.image_to_string(variant_img, config=config)
            txt = txt.replace("\x0c", "").strip("\n")
            passes.append(OCRPass(variant=variant_name, psm=psm, text=txt))
    return passes


_DATE_RE_YYYY_MM_DD = re.compile(r"\b(20\d{2})[-/.](\d{1,2})[-/.](\d{1,2})\b")
_DATE_RE_DD_MM_YYYY = re.compile(r"\b(\d{1,2})[-/.](\d{1,2})[-/.](20\d{2})\b")
_DATE_RE_DD_MM_YY = re.compile(r"\b(\d{1,2})[-/.](\d{1,2})[-/.](\d{2})\b")
_DATE_RE_DD_MON_YYYY = re.compile(
    r"\b(\d{1,2})\s*(JAN|FEB|MAR|APR|MAY|JUN|JUL|AUG|SEP|SEPT|OCT|NOV|DEC)[A-Z]*\s*(20\d{2})\b",
    re.I,
)
_DATE_RE_DD_MON_YY = re.compile(
    r"\b(\d{1,2})\s*(JAN|FEB|MAR|APR|MAY|JUN|JUL|AUG|SEP|SEPT|OCT|NOV|DEC)[A-Z]*\s*(\d{2})\b",
    re.I,
)
_DATE_RE_INVOICE_YYYYMMDD = re.compile(r"\b(20\d{2})(\d{2})(\d{2})\d{6,}\b")

_MONTH_MAP = {
    "JAN": 1,
    "FEB": 2,
    "MAR": 3,
    "APR": 4,
    "MAY": 5,
    "JUN": 6,
    "JUL": 7,
    "AUG": 8,
    "SEP": 9,
    "SEPT": 9,
    "OCT": 10,
    "NOV": 11,
    "DEC": 12,
}


def _try_make_date(y: int, m: int, d: int) -> date | None:
    try:
        return date(y, m, d)
    except ValueError:
        return None


def _extract_date_candidates(text: str) -> list[dict[str, Any]]:
    up = text.upper()
    lines = [ln.strip() for ln in text.splitlines() if ln.strip()]
    candidates: list[dict[str, Any]] = []

    def add(dt: date, raw: str, line: str, label: str | None, score: float, line_index: int) -> None:
        candidates.append(
            {
                "value": dt.isoformat(),
                "raw": raw,
                "line": line,
                "label": label,
                "score": score,
                "line_index": line_index,
            }
        )

    for i, ln in enumerate(lines):
        ln_up = ln.upper()
        prev_up = lines[i - 1].upper() if i > 0 else ""
        has_date_context = ("DATE" in ln_up) or ("DATE" in prev_up)
        label = "DATE" if has_date_context else None

        # Some receipts embed yyyymmdd inside invoice numbers.
        if ("INVOICE" in ln_up or "INV" in ln_up) and any(ch.isdigit() for ch in ln_up):
            for m in _DATE_RE_INVOICE_YYYYMMDD.finditer(ln_up):
                y, mm, dd = int(m.group(1)), int(m.group(2)), int(m.group(3))
                dt = _try_make_date(y, mm, dd)
                if dt:
                    score = 9.0 + (6.0 if has_date_context else 0.0)
                    add(dt, m.group(0), ln, label, score, i)

        # Year-first numeric
        for m in _DATE_RE_YYYY_MM_DD.finditer(ln):
            y, mm, dd = int(m.group(1)), int(m.group(2)), int(m.group(3))
            dt = _try_make_date(y, mm, dd)
            if dt:
                score = 10.0 + (6.0 if has_date_context else 0.0)
                if ":" in ln:
                    score += 0.5
                add(dt, m.group(0), ln, label, score, i)

        # Day-first 4-digit year
        for m in _DATE_RE_DD_MM_YYYY.finditer(ln):
            dd, mm, y = int(m.group(1)), int(m.group(2)), int(m.group(3))
            dt = _try_make_date(y, mm, dd)
            if dt is None and dd <= 12 and mm <= 12:
                dt = _try_make_date(y, dd, mm)
            if dt:
                score = 8.0 + (6.0 if has_date_context else 0.0)
                if ":" in ln:
                    score += 0.5
                add(dt, m.group(0), ln, label, score, i)

        # Day-first 2-digit year (assume 20xx)
        for m in _DATE_RE_DD_MM_YY.finditer(ln):
            dd, mm, yy = int(m.group(1)), int(m.group(2)), int(m.group(3))
            y = 2000 + yy
            dt = _try_make_date(y, mm, dd)
            if dt is None:
                dt = _try_make_date(y, dd, mm)
            if dt:
                score = 6.0 + (6.0 if has_date_context else 0.0)
                if ":" in ln:
                    score += 0.5
                add(dt, m.group(0), ln, label, score, i)

        # Month names
        for m in _DATE_RE_DD_MON_YYYY.finditer(ln):
            dd, mon, y = int(m.group(1)), m.group(2).upper(), int(m.group(3))
            mm = _MONTH_MAP.get(mon)
            if mm:
                dt = _try_make_date(y, mm, dd)
                if dt:
                    score = 9.0 + (6.0 if has_date_context else 0.0)
                    add(dt, m.group(0), ln, label, score, i)

        for m in _DATE_RE_DD_MON_YY.finditer(ln):
            dd, mon, yy = int(m.group(1)), m.group(2).upper(), int(m.group(3))
            mm = _MONTH_MAP.get(mon)
            y = 2000 + yy
            if mm:
                dt = _try_make_date(y, mm, dd)
                if dt:
                    score = 7.0 + (6.0 if has_date_context else 0.0)
                    add(dt, m.group(0), ln, label, score, i)

    if candidates and "DATE" not in up:
        uniq = {c["value"] for c in candidates}
        if len(uniq) == 1:
            for c in candidates:
                c["score"] += 1.0

    return candidates


_AMOUNT_RE = re.compile(
    r"(?<!\d)(\d{1,3}(?:,\d{3})+(?:\.\d{1,2})?|\d+\.\d{1,2}|\d+,\d{2}|\d+)(?!\d)"
)


def _norm_amount(raw_num: str) -> str | None:
    s = raw_num.strip()
    if "," in s and "." not in s and re.fullmatch(r"\d+,\d{2}", s):
        s = s.replace(",", ".")
    else:
        s = s.replace(",", "")
    try:
        d = Decimal(s)
    except InvalidOperation:
        return None
    d = d.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    if d < 0:
        return None
    return f"{d:.2f}"


def _line_has_non_payable_cue(ln_up: str) -> bool:
    return any(cue in ln_up for cue in NON_PAYABLE_CUES)


def _keyword_priority(ln_up: str) -> tuple[int, str | None]:
    for idx, kw in enumerate(PAYABLE_KEYWORDS_PRIORITY):
        if kw in ln_up:
            return (idx, kw)
    return (10_000, None)


def _extract_amount_candidates(text: str) -> list[dict[str, Any]]:
    lines = [ln.strip() for ln in text.splitlines()]
    candidates: list[dict[str, Any]] = []

    for i, ln in enumerate(lines):
        if not ln:
            continue

        ln_norm = re.sub(r"\s+", " ", ln).strip()
        ln_up = ln_norm.upper()
        pri, kw = _keyword_priority(ln_up)
        if kw is None:
            continue

        # Exclude clear non-payable totals.
        if "SUBTOTAL" in ln_up or "SUB TOTAL" in ln_up:
            continue

        in_tax_summary = (
            "GST SUMMARY" in ln_up
            or "TAX SUMMARY" in ln_up
            or (("GST" in ln_up or "TAX" in ln_up) and "SUMMARY" in ln_up)
        )
        if in_tax_summary and kw in ("TOTAL", "AMOUNT"):
            continue

        if _line_has_non_payable_cue(ln_up) and kw in ("TOTAL", "AMOUNT"):
            # Allow final-total lines like "Total Inclusive GST".
            if not any(tok in ln_up for tok in ("INCL", "INCLUSIVE", "INCLUDES", "AFTER", "ROUND", "ROUNDED", "ADJ")):
                continue

        amts: list[tuple[str, str]] = []
        for m in _AMOUNT_RE.finditer(ln_norm):
            raw = m.group(0)
            norm = _norm_amount(raw)
            if not norm:
                continue

            is_currency = ("RM" in ln_up) or ("$" in ln_up)
            is_strong_kw = kw in ("GRAND TOTAL", "TOTAL RM", "TOTAL: RM", "TOTAL AMOUNT", "TOTAL DUE", "AMOUNT DUE", "BALANCE DUE")

            if "." not in raw and not re.search(r",\d{2}$", raw) and not is_currency and not is_strong_kw:
                continue

            amts.append((raw, norm))

        if amts:
            raw, norm = amts[-1]

            desc_bonus = 0.0
            if any(tok in ln_up for tok in ("AFTER", "ADJ")):
                desc_bonus += 70.0
            if any(tok in ln_up for tok in ("ROUND", "ROUNDED", "ROUNDING")):
                desc_bonus += 60.0
            if any(tok in ln_up for tok in ("INCL", "INCLUSIVE", "INCLUDES")):
                desc_bonus += 45.0
            if "TAKEOUT" in ln_up:
                desc_bonus += 25.0
            if "RM" in ln_up:
                desc_bonus += 10.0
            if kw == "GRAND TOTAL":
                desc_bonus += 120.0

            score = 1000.0 - pri * 80.0 + desc_bonus + i * 0.02
            candidates.append(
                {
                    "value": norm,
                    "raw": raw,
                    "line": ln_norm,
                    "label": kw,
                    "score": score,
                    "line_index": i,
                }
            )
            continue

        # Split label/amount across lines: label-only line then amount-only next line
        if i + 1 < len(lines):
            nxt = lines[i + 1].strip()
            if not nxt:
                continue
            nxt_norm = re.sub(r"\s+", " ", nxt).strip()
            nxt_up = nxt_norm.upper()
            if _keyword_priority(nxt_up)[1] is not None:
                continue
            if _line_has_non_payable_cue(nxt_up):
                continue

            nxt_amts: list[tuple[str, str]] = []
            for m in _AMOUNT_RE.finditer(nxt_norm):
                raw = m.group(0)
                norm = _norm_amount(raw)
                if norm:
                    nxt_amts.append((raw, norm))

            if len(nxt_amts) == 1:
                raw, norm = nxt_amts[0]
                score = 900.0 - pri * 80.0 + i * 0.02
                candidates.append(
                    {
                        "value": norm,
                        "raw": raw,
                        "line": f"{ln_norm} / {nxt_norm}",
                        "label": kw,
                        "score": score,
                        "line_index": i,
                    }
                )

    return candidates


def _amount_support_points(text: str) -> dict[str, float]:
    lines = [ln.strip() for ln in text.splitlines() if ln.strip()]
    pts: dict[str, float] = {}

    cash: str | None = None
    change: str | None = None
    rounding_seen = False

    for ln in lines:
        ln_norm = re.sub(r"\s+", " ", ln).strip()
        ln_up = ln_norm.upper()

        if "ROUND" in ln_up:
            rounding_seen = True

        vals: list[str] = []
        for m in _AMOUNT_RE.finditer(ln_norm):
            raw = m.group(0)
            if "." not in raw and not re.search(r",\d{2}$", raw) and "RM" not in ln_up and "$" not in ln_up:
                continue
            norm = _norm_amount(raw)
            if norm:
                vals.append(norm)

        if not vals:
            continue

        if "PAYMENT" in ln_up:
            for v in vals:
                pts[v] = pts.get(v, 0.0) + 30.0

        if "TOTAL" in ln_up and any(tok in ln_up for tok in ("INCL", "INCLUSIVE", "INCLUDES")):
            for v in vals:
                pts[v] = pts.get(v, 0.0) + 15.0

        if "TOTAL" in ln_up and any(tok in ln_up for tok in ("AFTER", "ROUND", "ROUNDED", "ADJ")):
            for v in vals:
                pts[v] = pts.get(v, 0.0) + 20.0

        if rounding_seen and "TOTAL" in ln_up and not any(tok in ln_up for tok in ("AFTER", "ROUND", "ROUNDED", "ROUNDING", "ADJ")):
            for v in vals:
                pts[v] = pts.get(v, 0.0) + 12.0

        if ln_up.startswith("TOTAL") and ":" in ln_up[:8]:
            for v in vals:
                pts[v] = pts.get(v, 0.0) + 30.0


        if "CASH" in ln_up and "CASHIER" not in ln_up and "CASHBACK" not in ln_up:
            cash = vals[-1]

        if "CHANGE" in ln_up:
            change = vals[-1]

    if cash is not None and change is not None:
        try:
            total = (Decimal(cash) - Decimal(change)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
            if total >= 0:
                norm = f"{total:.2f}"
                pts[norm] = pts.get(norm, 0.0) + 40.0
        except Exception:
            pass

    return pts


def _best_candidate(cands: list[dict[str, Any]]) -> dict[str, Any] | None:
    if not cands:
        return None
    return sorted(cands, key=lambda c: (-float(c["score"]), -int(c.get("line_index", 0))))[0]


def select_with_field_evidence(observations: list[OCRPass], field: str) -> dict[str, Any]:
    assert field in FIELDS

    def variant_weight(variant: str, psm: int) -> float:
        if variant in {"gray_autocontrast", "invert_autocontrast"}:
            v = 1.0
        elif variant == "upscale2x":
            v = 0.95
        elif variant == "contrast_1.6":
            v = 0.85
        elif variant == "sharpen":
            v = 0.8
        elif variant.startswith("threshold_"):
            v = 0.65
        else:
            v = 0.75

        if psm == 6:
            p = 1.0
        elif psm == 4:
            p = 0.9
        else:
            p = 0.75

        return v * p

    all_candidates: list[dict[str, Any]] = []
    support: dict[str, float] = {}
    votes: dict[str, float] = {}

    for obs in observations:
        if not obs.text.strip():
            continue

        if field == "total_amount":
            sp = _amount_support_points(obs.text)
            for k, v in sp.items():
                support[k] = support.get(k, 0.0) + float(v)

        cands = _extract_date_candidates(obs.text) if field == "date" else _extract_amount_candidates(obs.text)
        if cands:
            top = max(cands, key=lambda c: float(c.get("score", 0.0)))
            votes[top["value"]] = votes.get(top["value"], 0.0) + variant_weight(obs.variant, obs.psm)

        for c in cands:
            all_candidates.append(
                {
                    "value": c["value"],
                    "evidence": {
                        "raw_text": c.get("line"),
                        "label": c.get("label"),
                        "bbox": None,
                        "ocr_source": f"{obs.variant};psm={obs.psm}",
                    },
                    "score": float(c.get("score", 0.0)),
                    "line_index": int(c.get("line_index", 0)),
                }
            )

    rec = _default_field_record()
    rec["review_reason"] = None

    if not all_candidates:
        rec["state"] = "unresolved"
        rec["value"] = None
        rec["review_reason"] = f"no {field} candidates found"
        rec["candidates"] = []
        return rec

    # Keep all candidates (dedupe by value+source+raw_text)
    seen = set()
    c_list = []
    for c in sorted(all_candidates, key=lambda x: (-float(x.get("score", 0.0)), -int(x.get("line_index", 0)))):
        key = (c["value"], c["evidence"]["ocr_source"], c["evidence"]["raw_text"])
        if key in seen:
            continue
        seen.add(key)
        c_list.append(
            {
                "value": c["value"],
                "raw": c["evidence"]["raw_text"],
                "label": c["evidence"]["label"],
                "bbox": None,
                "ocr_source": c["evidence"]["ocr_source"],
            }
        )
    rec["candidates"] = c_list

    stats: dict[str, dict[str, Any]] = {}
    for c in all_candidates:
        v = c["value"]
        st = stats.setdefault(v, {"best": c, "best_score": -1.0})
        sc = float(c.get("score", 0.0))
        if sc > float(st["best_score"]):
            st["best_score"] = sc
            st["best"] = c


    preferred_values: set[str] | None = None
    if field == "date":
        trusted_sources = {
            "gray_autocontrast;psm=6",
            "invert_autocontrast;psm=6",
            "gray_autocontrast;psm=4",
            "invert_autocontrast;psm=4",
        }
        trusted = [
            c
            for c in all_candidates
            if c["evidence"]["ocr_source"] in trusted_sources
            and (c["evidence"].get("label") or "").upper() == "DATE"
        ]
        if trusted:
            counts: dict[str, int] = {}
            for c in trusted:
                counts[c["value"]] = counts.get(c["value"], 0) + 1
            mx = max(counts.values())
            preferred_values = {v for v, cnt in counts.items() if cnt == mx}

    def _date_label_strength(label: str | None) -> int:
        return 50 if (label or "").upper() == "DATE" else 0

    def _total_label_strength(label: str | None) -> int:
        lab = (label or "").upper()
        if lab == "GRAND TOTAL":
            return 100
        if lab in {"TOTAL RM", "TOTAL: RM"}:
            return 95
        if lab == "TOTAL AMOUNT":
            return 90
        if lab in {"TOTAL DUE", "AMOUNT DUE", "BALANCE DUE"}:
            return 85
        if lab in {"NETT TOTAL", "NET TOTAL"}:
            return 80
        if lab == "TOTAL":
            return 55
        if lab == "AMOUNT":
            return 45
        return 0

    best_value: str | None = None
    best_entry: dict[str, Any] | None = None
    best_rank: tuple[Any, ...] | None = None

    for v, st in stats.items():
        if preferred_values is not None and v not in preferred_values:
            continue
        entry = st["best"]
        label = entry["evidence"].get("label")
        vote_sum = float(votes.get(v, 0.0))

        if field == "date":
            rank = (_date_label_strength(label), vote_sum, float(st["best_score"]))
        else:
            rank = (float(support.get(v, 0.0)), _total_label_strength(label), vote_sum, float(st["best_score"]))

        if best_rank is None or rank > best_rank:
            best_rank = rank
            best_value = v
            best_entry = entry

    assert best_value is not None and best_entry is not None

    rec["value"] = best_value
    rec["state"] = "candidate"
    rec["evidence"] = best_entry["evidence"]

    if field == "date":
        lbl = best_entry["evidence"].get("label")
        if _date_label_strength(lbl) > 0:
            rec["state"] = "confirmed"
            rec["confirmation_reason"] = f"Date selected with DATE context (vote={votes.get(best_value, 0.0):.2f})"
        elif len(stats) == 1:
            rec["state"] = "confirmed"
            rec["confirmation_reason"] = "Only one unique date found across OCR passes"
        elif float(votes.get(best_value, 0.0)) >= 2.0:
            rec["state"] = "confirmed"
            rec["confirmation_reason"] = "Date selected by weighted vote across OCR passes"
        else:
            rec["review_reason"] = "date not explicitly labeled"

    if field == "total_amount":
        lbl = best_entry["evidence"].get("label")
        ls = _total_label_strength(lbl)
        sp = float(support.get(best_value, 0.0))
        if ls >= 80 or sp >= 40.0 or float(votes.get(best_value, 0.0)) >= 2.0:
            rec["state"] = "confirmed"
            rec["confirmation_reason"] = f"Total selected using label/support evidence (support={sp:.1f})"
        else:
            rec["review_reason"] = "total label too generic"

    return rec


def copy_field_record(r: dict[str, Any]) -> dict[str, Any]:
    return json.loads(json.dumps(r))


def protect_confirmed_field(before: dict[str, Any], proposal: dict[str, Any]) -> dict[str, Any]:
    def label_strength(label: str | None) -> int:
        lab = (label or "").upper()
        if lab == "DATE":
            return 100
        if lab == "GRAND TOTAL":
            return 95
        if lab in {"TOTAL RM", "TOTAL: RM"}:
            return 90
        if lab == "TOTAL AMOUNT":
            return 85
        if lab in {"TOTAL DUE", "AMOUNT DUE", "BALANCE DUE"}:
            return 80
        if lab in {"NETT TOTAL", "NET TOTAL"}:
            return 75
        if lab == "TOTAL":
            return 55
        if lab == "AMOUNT":
            return 45
        return 0

    def merge_candidates(dst: dict[str, Any], src: dict[str, Any]) -> None:
        if not src.get("candidates"):
            return
        dst.setdefault("candidates", [])
        existing = {(c.get("value"), c.get("ocr_source"), c.get("raw")) for c in dst.get("candidates", [])}
        for c in src.get("candidates", []):
            key = (c.get("value"), c.get("ocr_source"), c.get("raw"))
            if key not in existing:
                dst["candidates"].append(c)
                existing.add(key)

    before_v = before.get("value")
    prop_v = proposal.get("value")

    if before.get("state") == "confirmed":
        kept = copy_field_record(before)
        merge_candidates(kept, proposal)
        if prop_v is not None and prop_v != before_v:
            kept["review_reason"] = f"contradiction: proposed {prop_v} from {proposal.get('evidence', {}).get('ocr_source')}"
        return kept

    if prop_v is None and before_v is not None:
        kept = copy_field_record(before)
        merge_candidates(kept, proposal)
        return kept

    if before_v is None and prop_v is not None:
        kept = copy_field_record(proposal)
        merge_candidates(kept, before)
        if before.get("changes"):
            kept["changes"] = before["changes"]
        return kept

    if before_v is None and prop_v is None:
        kept = copy_field_record(proposal)
        merge_candidates(kept, before)
        if before.get("changes"):
            kept["changes"] = before["changes"]
        return kept

    assert before_v is not None and prop_v is not None

    if before_v == prop_v:
        kept = copy_field_record(before)
        merge_candidates(kept, proposal)
        return kept

    before_ls = label_strength(before.get("evidence", {}).get("label"))
    prop_ls = label_strength(proposal.get("evidence", {}).get("label"))

    if proposal.get("state") == "confirmed" and before.get("state") != "confirmed":
        kept = copy_field_record(proposal)
        merge_candidates(kept, before)
        if before.get("changes"):
            kept["changes"] = before["changes"]
        return kept

    if prop_ls > before_ls:
        kept = copy_field_record(proposal)
        merge_candidates(kept, before)
        if before.get("changes"):
            kept["changes"] = before["changes"]
        return kept

    kept = copy_field_record(before)
    merge_candidates(kept, proposal)
    kept["review_reason"] = kept.get("review_reason") or f"contradiction: proposed {prop_v}"
    return kept


def first_pass(records: dict[str, Any], images: list[Path]) -> None:
    for image_path in images:
        pending = [field for field in FIELDS if not records[image_path.name][field]["first_pass_done"]]
        if not pending:
            continue

        print(f"first pass: {image_path.name}; fields={pending}", flush=True)
        observations = recognize_with_text_and_regions(image_path)

        for field in pending:
            before = copy_field_record(records[image_path.name][field])
            proposal = select_with_field_evidence(observations, field)
            proposal["first_pass_done"] = True
            if proposal.get("review_reason") is None:
                proposal["review_reason"] = ""
            records[image_path.name][field] = protect_confirmed_field(before, proposal)

        save_working_records(records)
        print(f"checkpoint saved: {image_path.name}", flush=True)


def first_pass_completion_check(records: dict[str, Any], images: list[Path]) -> None:
    with LEDGER_PATH.open("r", encoding="utf-8") as f:
        reopened = json.load(f)

    missing = []
    counts: dict[str, int] = {}
    for p in images:
        for field in FIELDS:
            r = reopened[p.name][field]
            if not r.get("first_pass_done"):
                missing.append((p.name, field))
            counts[r.get("state", "unknown")] = counts.get(r.get("state", "unknown"), 0) + 1

    print(f"first pass check: states={counts} missing_first_pass={len(missing)}", flush=True)
    if missing:
        raise RuntimeError(f"first pass incomplete: {missing[:5]} ...")


def fields_requiring_review(records: dict[str, Any]) -> list[dict[str, Any]]:
    queue: list[dict[str, Any]] = []
    for filename in sorted(records.keys()):
        for field in FIELDS:
            r = records[filename][field]
            reason = None
            if r.get("state") != "confirmed":
                reason = r.get("review_reason") or "not confirmed"
            elif r.get("review_reason") and str(r.get("review_reason")).startswith("contradiction:"):
                reason = r.get("review_reason")

            if reason:
                queue.append(
                    {
                        "filename": filename,
                        "field": field,
                        "reason": reason,
                        "status": "pending",
                        "outcome": None,
                    }
                )
    return queue


def save_review_queue(queue: list[dict[str, Any]]) -> None:
    with QUEUE_PATH.open("w", encoding="utf-8") as f:
        json.dump(queue, f, indent=2, ensure_ascii=False)


def _review_crop_bbox(image_path: Path, field: str) -> tuple[int, int, int, int]:
    img = Image.open(image_path)
    w, h = img.size
    if field == "date":
        return (0, 0, w, int(h * 0.35))
    return (0, int(h * 0.55), w, h)


def review_original_field_region(image_path: Path, field: str, before: dict[str, Any]) -> dict[str, Any]:
    bbox = _review_crop_bbox(image_path, field)
    observations = recognize_with_text_and_regions(image_path, crop_bbox=bbox)
    proposal = select_with_field_evidence(observations, field)

    # Inject bbox into evidence for traceability
    if proposal.get("evidence"):
        proposal["evidence"]["bbox"] = list(bbox)
        src = proposal["evidence"].get("ocr_source")
        proposal["evidence"]["ocr_source"] = f"review_crop={field}:{bbox};{src}"

    # If this review can confirm (e.g. label shows up), upgrade state
    proposal["first_pass_done"] = True

    if proposal.get("state") == "confirmed":
        proposal["review_reason"] = None
    else:
        proposal["review_reason"] = proposal.get("review_reason") or "review did not confirm"

    return proposal


def run_reviews(records: dict[str, Any], images_by_name: dict[str, Path]) -> list[dict[str, Any]]:
    queue = fields_requiring_review(records)
    save_review_queue(queue)

    for entry in queue:
        if entry["status"] != "pending":
            continue

        filename = entry["filename"]
        field = entry["field"]
        before = copy_field_record(records[filename][field])

        print(f"local review: {filename}; field={field}; reason={entry['reason']}", flush=True)
        proposal = review_original_field_region(images_by_name[filename], field, before)

        updated = protect_confirmed_field(before, proposal)
        records[filename][field] = updated
        save_working_records(records)

        # Update queue entry
        entry["outcome"] = {
            "proposed": proposal.get("value"),
            "state": proposal.get("state"),
            "source": proposal.get("evidence", {}).get("ocr_source"),
        }
        if records[filename][field].get("state") == "confirmed":
            entry["status"] = "reviewed"
        else:
            entry["status"] = "unresolved"

        save_review_queue(queue)

    return queue


def materialize_prescribed_rows(records: dict[str, Any]) -> list[list[Any]]:
    rows: list[list[Any]] = []
    for filename in sorted(records.keys()):
        row = [filename]
        for field in FIELDS:
            r = records[filename][field]
            v = r.get("value")
            row.append(v)
        rows.append(row)
    return rows


def write_required_workbook(rows: list[list[Any]]) -> None:
    wb = Workbook()
    # Remove default
    wb.remove(wb.active)

    ws = wb.create_sheet("results")
    ws.append(["filename", "date", "total_amount"])
    for r in rows:
        ws.append(r)

    wb.save(OUT_XLSX)


def reopen_required_workbook() -> list[list[Any]]:
    wb = load_workbook(OUT_XLSX, data_only=True)
    if wb.sheetnames != ["results"]:
        raise RuntimeError(f"Workbook must contain only 'results' sheet, found: {wb.sheetnames}")

    ws = wb["results"]
    out: list[list[Any]] = []
    for row in ws.iter_rows(values_only=True):
        out.append(list(row))
    return out


def sorted_input_coverage_is_complete(reopened_rows: list[list[Any]], images: list[Path]) -> bool:
    if not reopened_rows:
        return False
    data_rows = reopened_rows[1:]
    if len(data_rows) != len(images):
        return False

    expected = [p.name for p in sorted(images, key=lambda p: p.name)]
    actual = [r[0] for r in data_rows]
    return expected == actual


def report_counts(records: dict[str, Any], queue: list[dict[str, Any]]) -> None:
    counts = {"confirmed": 0, "candidate": 0, "unresolved": 0}
    for fn in records:
        for field in FIELDS:
            st = records[fn][field].get("state")
            counts[st] = counts.get(st, 0) + 1

    q_counts = {"pending": 0, "reviewed": 0, "unresolved": 0}
    for e in queue:
        q_counts[e["status"]] = q_counts.get(e["status"], 0) + 1

    print(f"final ledger counts: {counts}; queue counts: {q_counts}", flush=True)


def main() -> None:
    images = enumerate_input_images_sorted()
    if not images:
        raise RuntimeError(f"No images found under {IMG_DIR}")

    records = load_working_records_or_initialize(images)
    images_by_name = {p.name: p for p in images}

    first_pass(records, images)
    first_pass_completion_check(records, images)

    queue = run_reviews(records, images_by_name)

    rows = materialize_prescribed_rows(records)
    write_required_workbook(rows)

    reopened = reopen_required_workbook()
    reopened_rows = reopened[1:]
    if reopened[0] != ["filename", "date", "total_amount"]:
        raise RuntimeError(f"Header mismatch: {reopened[0]}")

    if [r for r in reopened_rows] != rows:
        raise RuntimeError("Reopened workbook rows do not match written rows")

    if not sorted_input_coverage_is_complete(reopened, images):
        raise RuntimeError("Workbook input coverage/order is incomplete")

    report_counts(records, queue)


if __name__ == "__main__":
    main()
