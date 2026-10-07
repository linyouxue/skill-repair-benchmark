import pandas as pd
from openpyxl import load_workbook

wb = load_workbook("test-supply.xlsx")

# 1. PWT Data
pwt = pd.read_stata("pwt1001.xlsx")
pwt_geo = pwt[pwt['country'] == 'Georgia'].dropna(subset=['rnna'])
pwt_dict = dict(zip(pwt_geo['year'], pwt_geo['rnna']))

pwt_sheet = wb['PWT']
# row 2 is index 1, value is year 1990
for row in range(2, pwt_sheet.max_row + 1):
    year_val = pwt_sheet.cell(row=row, column=1).value
    if year_val in pwt_dict:
        # Fill Column B (which is index 2)
        pwt_sheet.cell(row=row, column=2).value = pwt_dict[year_val]

# 2. WEO Data
weo = pd.read_csv("geo_weo.csv")
weo_r_dict = dict(zip(weo['year'], weo['NGDP_R']))
weo_rpch_dict = dict(zip(weo['year'], weo['NGDP_RPCH']))

weo_sheet = wb['WEO_Data']
# Years start from row 8. Column B has years, Column C is NGDP_R, Column D is NGDP_RPCH
last_known_growth = weo_rpch_dict[2027] / 100.0 if 2027 in weo_rpch_dict else 0.045
weo_rpch_dict[2027] = last_known_growth * 100

for row in range(8, 52): # Up to 2043 is row 51
    year_val = weo_sheet.cell(row=row, column=2).value
    if year_val is None:
        continue
    year_val = int(year_val)
    
    if year_val <= 2027:
        if year_val in weo_r_dict:
            weo_sheet.cell(row=row, column=3).value = weo_r_dict[year_val]
        if year_val in weo_rpch_dict:
            weo_sheet.cell(row=row, column=4).value = weo_rpch_dict[year_val]
    else:
        # Extend from 2028 to 2043
        # In Excel we can write formulas to extend NGDP_RPCH (Column D) and calculate NGDP_R (Column C)
        # Assuming we need to keep formulas!
        weo_sheet.cell(row=row, column=4).value = f"=D{row-1}"
        weo_sheet.cell(row=row, column=3).value = f"=C{row-1}*(1+D{row}/100)"

# 3. ECB Data
ecb = pd.read_csv("ecb_data.csv")
ecb_dict = dict(zip(ecb['TIME_PERIOD'], ecb['OBS_VALUE']))

cfc_sheet = wb['CFC data']
# Let's inspect CFC data structure. row 2 is 1996
for row in range(2, cfc_sheet.max_row + 1):
    year_val = cfc_sheet.cell(row=row, column=2).value
    if year_val in ecb_dict:
        cfc_sheet.cell(row=row, column=3).value = ecb_dict[year_val]

# "Link the capital stock data from "PWT", and calculate the depreciation rate using column C&D"
# PWT Data has years 1990 to 2019 (pwt 10.01 ends 2019). We need to link it.
# Wait, column D is Capital Stock in 'CFC data'. 
for row in range(2, cfc_sheet.max_row + 1):
    year_val = cfc_sheet.cell(row=row, column=2).value
    if year_val is None:
        continue
    # Link to PWT data
    # In PWT sheet, year 1996 is row 8 (1990 is row 2)
    pwt_row = int(year_val) - 1990 + 2
    cfc_sheet.cell(row=row, column=4).value = f"='PWT'!B{pwt_row}"
    # Calculate depreciation rate using C & D: Rate = CFC / Capital Stock = C / D
    cfc_sheet.cell(row=row, column=5).value = f"=C{row}/D{row}"

# Save temp
wb.save("temp.xlsx")
print("Saved to temp.xlsx")
