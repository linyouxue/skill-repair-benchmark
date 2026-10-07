from openpyxl import load_workbook
import pandas as pd
import numpy as np

file_path = '/root/output/rar_result.xlsx'
wb = load_workbook(file_path)
ws_answer = wb['Answer']
ws_gold = wb['Gold price']

gold_2025 = []
for i in range(1, 10):
    row_idx = 421 + i
    val = ws_gold.cell(row=row_idx, column=2).value
    if val is not None:
        gold_2025.append(float(val))
gold_price_2025_avg = np.mean(gold_2025)

for row in range(1, 30):
    for col in range(3, 27):
        ws_answer.cell(row=row, column=col).value = None

max_row = 430
ws_answer['C3'] = "=NORM.S.INV(0.95)"
ws_answer['C4'] = f"='Gold price'!D{max_row}"
ws_answer['C5'] = "=C4*SQRT(12/3)"
ws_answer['C6'] = f"='Gold price'!E{max_row}"

df_value = pd.read_excel(file_path, sheet_name='Value', header=7)
df_volume = pd.read_excel(file_path, sheet_name='Volume', header=7)

row_2025_idx = None
for idx, val in enumerate(df_value.iloc[:, 0]):
    if str(val) == '2025':
        row_2025_idx = idx
        break

countries = {}
if row_2025_idx is not None:
    for col in df_value.columns[2:]:
        cname = col.split(':')[0].strip()
        val = df_value.at[row_2025_idx, col]
        if pd.notna(val) and str(val).strip() != '':
            countries[cname] = float(val)

row_2025_idx_vol = None
for idx, val in enumerate(df_volume.iloc[:, 0]):
    if str(val) == '2025':
        row_2025_idx_vol = idx
        break
if row_2025_idx_vol is not None:
    for col in df_volume.columns[2:]:
        cname = col.split(':')[0].strip()
        val = df_volume.at[row_2025_idx_vol, col]
        if pd.notna(val) and str(val).strip() != '':
            if cname not in countries:
                vol = float(val)
                if 'Thous' in col or 'Thous.' in col:
                    vol = vol / 1000.0
                countries[cname] = vol * gold_price_2025_avg

def get_col_letter(col_num):
    res = ""
    while col_num > 0:
        col_num, rem = divmod(col_num - 1, 26)
        res = chr(65 + rem) + res
    return res

start_col = 3
for i, (country, value) in enumerate(countries.items()):
    col_idx = start_col + i
    col_letter = get_col_letter(col_idx)
    ws_answer[f'{col_letter}11'] = country
    ws_answer[f'{col_letter}12'] = value
    ws_answer[f'{col_letter}13'] = f"={col_letter}12*$C$3*$C$5/100"

df_total = pd.read_excel(file_path, sheet_name='Total Reserves', header=7)
row_2025_idx_tot = None
for idx, val in enumerate(df_total.iloc[:, 0]):
    if str(val) == '2025':
        row_2025_idx_tot = idx
        break

total_reserves = {}
if row_2025_idx_tot is not None:
    for col in df_total.columns[2:]:
        cname = col.split(':')[0].strip()
        val = df_total.at[row_2025_idx_tot, col]
        if pd.notna(val) and str(val).strip() != '':
            total_reserves[cname] = float(val)

curr_col = 3
for i, (country, value) in enumerate(countries.items()):
    if country in total_reserves:
        col_letter = get_col_letter(curr_col)
        step2_col_letter = get_col_letter(start_col + i)
        
        ws_answer[f'{col_letter}20'] = country
        ws_answer[f'{col_letter}21'] = f"={step2_col_letter}12"
        ws_answer[f'{col_letter}22'] = f"={step2_col_letter}13"
        formula = f'=XLOOKUP({col_letter}20&"*", \'Total Reserves\'!$C$8:$R$8, \'Total Reserves\'!$C$17:$R$17, "", 2)'
        ws_answer[f'{col_letter}23'] = formula
        ws_answer[f'{col_letter}24'] = f"={col_letter}22/{col_letter}23*100"
        curr_col += 1

wb.save(file_path)
print("Done fixing all")
