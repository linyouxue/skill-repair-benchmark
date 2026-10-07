from openpyxl import load_workbook
file_path = '/root/output/rar_result.xlsx'
wb = load_workbook(file_path)
ws = wb['Answer']
ws['B2'] = None
ws['B3'] = "Z-score (95% one-sided)"
ws['C3'] = "=NORM.S.INV(0.95)"
# Make sure C2 is clear
ws['C2'] = None
wb.save(file_path)
