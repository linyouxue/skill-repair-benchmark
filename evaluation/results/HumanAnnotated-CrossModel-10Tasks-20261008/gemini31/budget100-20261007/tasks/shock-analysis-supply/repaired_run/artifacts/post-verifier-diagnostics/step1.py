import pandas as pd
from openpyxl import load_workbook
import urllib.request
import zipfile

# 1. PWT Data
pwt = pd.read_excel('pwt_data.xlsx')
georgia_pwt = pwt[pwt['country'] == 'Georgia']
georgia_pwt = georgia_pwt[['year', 'rnna']].dropna()
georgia_pwt['year'] = georgia_pwt['year'].astype(int)

wb = load_workbook('test-supply.xlsx')
ws_pwt = wb['PWT']

# Get rows for the years in the PWT sheet
for row in range(2, ws_pwt.max_row + 1):
    year_cell = ws_pwt.cell(row=row, column=1).value
    if year_cell is not None:
        try:
            year = int(year_cell)
            val = georgia_pwt[georgia_pwt['year'] == year]['rnna'].values
            if len(val) > 0:
                ws_pwt.cell(row=row, column=2).value = float(val[0])
        except ValueError:
            pass

wb.save('test-supply.xlsx')
print("PWT data filled")
