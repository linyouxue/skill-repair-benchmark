import pandas as pd
from openpyxl import load_workbook
import requests

wb = load_workbook('test-supply.xlsx')
ws = wb['WEO_Data']

for row in range(2, 60):
    year_cell = ws.cell(row=row, column=2).value
    try:
        year = int(year_cell)
        if year == 2000:
            ws.cell(row=row, column=3).value = 22802.8359375
            break
    except (ValueError, TypeError):
        pass

wb.save('test-supply.xlsx')
