import pandas as pd
from openpyxl import load_workbook

wb = load_workbook("test-supply.xlsx")

# ==================== STEP 1: data collection ====================

# 1. PWT Data
pwt = pd.read_excel("pwt110.xlsx", sheet_name='Data')
pwt_geo = pwt[pwt['country'] == 'Georgia']
pwt_rnna = dict(zip(pwt_geo['year'], pwt_geo['rnna']))
# Wait, are there other columns we need? Prompt: "fill the relevant data in column B and C in sheet PWT"
# test-supply PWT sheet has year in A, rnna in B. Maybe C is something else? Let's check openpyxl directly
pwt_sheet = wb['PWT']
print("PWT cols:", pwt_sheet.cell(1, 1).value, pwt_sheet.cell(1, 2).value, pwt_sheet.cell(1, 3).value)
