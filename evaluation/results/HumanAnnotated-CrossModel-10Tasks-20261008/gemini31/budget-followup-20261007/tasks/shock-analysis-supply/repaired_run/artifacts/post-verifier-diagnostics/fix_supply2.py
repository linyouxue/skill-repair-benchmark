import pandas as pd
from openpyxl import load_workbook
import openpyxl

wb = openpyxl.load_workbook('test-supply.xlsx')
pwt_sheet = wb['PWT']
print("Row 1: ", [pwt_sheet.cell(1, i).value for i in range(1, 10)])
print("Row 2: ", [pwt_sheet.cell(2, i).value for i in range(1, 10)])
