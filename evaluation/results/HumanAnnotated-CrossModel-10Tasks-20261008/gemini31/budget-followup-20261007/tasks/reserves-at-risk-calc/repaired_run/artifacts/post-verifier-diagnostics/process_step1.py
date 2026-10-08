import pandas as pd
import yfinance as yf
from openpyxl import load_workbook
import numpy as np

# Load original excel
file_path = '/root/data/test-rar.xlsx'
wb = load_workbook(file_path)
ws_gold = wb['Gold price']

wb_data = pd.read_excel('cmo.xlsx', sheet_name='Monthly Prices', header=4)
wb_gold_df = wb_data[['Unnamed: 0', 'Gold']].copy()
wb_gold_df.columns = ['Period', 'GoldPrice']
wb_gold_df['Period'] = wb_gold_df['Period'].astype(str)
wb_gold_df = wb_gold_df[wb_gold_df['Period'].str.match(r'\d{4}M\d{2}')].copy()
wb_gold_df['Period'] = wb_gold_df['Period'].str.replace(r'M0', 'M', regex=True)
wb_gold_df = wb_gold_df.set_index('Period')

gold = yf.Ticker('GC=F')
hist = gold.history(period='max', interval='1d')
hist = hist[hist.index.year >= 2024]
monthly = hist.resample('ME').last()
monthly['Period'] = monthly.index.year.astype(str) + 'M' + monthly.index.month.astype(str)
monthly = monthly.set_index('Period')

df_imf = pd.read_excel(file_path, sheet_name='Gold price')
last_price = None

for r_idx, period in enumerate(df_imf['Unnamed: 0'], start=2):
    if period in wb_gold_df.index and pd.notna(wb_gold_df.loc[period, 'GoldPrice']):
        price = wb_gold_df.loc[period, 'GoldPrice']
    elif period in monthly.index and pd.notna(monthly.loc[period, 'Close']):
        price = monthly.loc[period, 'Close']
    else:
        price = last_price
    
    if price is not None:
        price = float(price)

    ws_gold.cell(row=r_idx, column=2, value=price)
    if price is not None:
        last_price = price

    if r_idx > 2:
        ws_gold.cell(row=r_idx, column=3, value=f"=LN(B{r_idx}/B{r_idx-1})*100")
        
    if r_idx >= 4:
        ws_gold.cell(row=r_idx, column=4, value=f"=STDEV.S(C{r_idx-2}:C{r_idx})")
        
    if r_idx >= 13:
        ws_gold.cell(row=r_idx, column=5, value=f"=STDEV.S(C{r_idx-11}:C{r_idx})")

ws_answer = wb['Answer']
ws_answer['B2'] = "=NORM.S.INV(0.95)"
# Find max row in Gold price sheet
max_row = len(df_imf) + 1 
ws_answer['B3'] = f"='Gold price'!D{max_row}"
ws_answer['B4'] = "=B3*SQRT(12/3)"
ws_answer['B5'] = f"='Gold price'!E{max_row}"

wb.save('/root/output/rar_result.xlsx')
print("Step 1 done")
