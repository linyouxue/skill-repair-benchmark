import pandas as pd
import yfinance as yf
from openpyxl import load_workbook
import numpy as np

# Load original excel
file_path = '/root/data/test-rar.xlsx'
wb = load_workbook(file_path)
ws_gold = wb['Gold price']
ws_answer = wb['Answer']

# ----------------- STEP 1 -----------------
wb_data = pd.read_excel('cmo.xlsx', sheet_name='Monthly Prices', header=4)
wb_gold_df = wb_data[['Unnamed: 0', 'Gold']].copy()
wb_gold_df.columns = ['Period', 'GoldPrice']
wb_gold_df['Period'] = wb_gold_df['Period'].astype(str)
wb_gold_df = wb_gold_df[wb_gold_df['Period'].str.match(r'\d{4}M\d{2}')].copy()
wb_gold_df['Period'] = wb_gold_df['Period'].str.replace(r'M0', 'M', regex=True)
wb_gold_df = wb_gold_df.set_index('Period')

gold = yf.Ticker('GC=F')
hist = gold.history(period='max', interval='1d')
hist = hist[hist.index.year >= 2024]
monthly = hist.resample('ME').last()
monthly['Period'] = monthly.index.year.astype(str) + 'M' + monthly.index.month.astype(str)
monthly = monthly.set_index('Period')

df_imf = pd.read_excel(file_path, sheet_name='Gold price')
last_price = None

for r_idx, period in enumerate(df_imf['Unnamed: 0'], start=2):
    if period in wb_gold_df.index and pd.notna(wb_gold_df.loc[period, 'GoldPrice']):
        price = wb_gold_df.loc[period, 'GoldPrice']
    elif period in monthly.index and pd.notna(monthly.loc[period, 'Close']):
        price = monthly.loc[period, 'Close']
    else:
        price = last_price
    
    if price is not None:
        price = float(price)

    ws_gold.cell(row=r_idx, column=2, value=price)
    if price is not None:
        last_price = price

    if r_idx > 2:
        ws_gold.cell(row=r_idx, column=3, value=f"=LN(B{r_idx}/B{r_idx-1})*100")
        
    if r_idx >= 4:
        ws_gold.cell(row=r_idx, column=4, value=f"=STDEV.S(C{r_idx-2}:C{r_idx})")
        
    if r_idx >= 13:
        ws_gold.cell(row=r_idx, column=5, value=f"=STDEV.S(C{r_idx-11}:C{r_idx})")

max_row = len(df_imf) + 1 
ws_answer['C2'] = "=NORM.S.INV(0.95)"
ws_answer['C3'] = f"='Gold price'!D{max_row}"
ws_answer['C4'] = "=C3*SQRT(12/3)"
ws_answer['C5'] = f"='Gold price'!E{max_row}"

# ----------------- STEP 2 -----------------
df_value = pd.read_excel(file_path, sheet_name='Value', header=7)
df_volume = pd.read_excel(file_path, sheet_name='Volume', header=7)

# We need to get the prices we just mapped
gold_prices_2025 = []
for period in [f'2025M{i}' for i in range(1, 10)]:
    if period in wb_gold_df.index and pd.notna(wb_gold_df.loc[period, 'GoldPrice']):
        gold_prices_2025.append(float(wb_gold_df.loc[period, 'GoldPrice']))
    elif period in monthly.index and pd.notna(monthly.loc[period, 'Close']):
        gold_prices_2025.append(float(monthly.loc[period, 'Close']))
gold_price_2025_avg = np.mean(gold_prices_2025)

row_2025_idx = None
for idx, val in enumerate(df_value.iloc[:, 0]):
    if str(val) == '2025':
        row_2025_idx = idx
        break

countries_with_value_2025 = {}
if row_2025_idx is not None:
    for col in df_value.columns[2:]:
        country_name = col.split(':')[0].strip()
        val_2025 = df_value.at[row_2025_idx, col]
        if pd.notna(val_2025) and str(val_2025).strip() != '':
            countries_with_value_2025[country_name] = float(val_2025)

row_2025_idx_vol = None
for idx, val in enumerate(df_volume.iloc[:, 0]):
    if str(val) == '2025':
        row_2025_idx_vol = idx
        break

if row_2025_idx_vol is not None:
    for col in df_volume.columns[2:]:
        country_name = col.split(':')[0].strip()
        val_2025 = df_volume.at[row_2025_idx_vol, col]
        if pd.notna(val_2025) and str(val_2025).strip() != '':
            if country_name not in countries_with_value_2025:
                vol = float(val_2025)
                if 'Thous' in col or 'Thous.' in col:
                    vol = vol / 1000.0
                value = vol * gold_price_2025_avg
                countries_with_value_2025[country_name] = value

start_col = 3 # Column C
def get_col_letter(col_num):
    res = ""
    while col_num > 0:
        col_num, rem = divmod(col_num - 1, 26)
        res = chr(65 + rem) + res
    return res

for i, (country, value) in enumerate(countries_with_value_2025.items()):
    col_idx = start_col + i
    col_letter = get_col_letter(col_idx)
    
    ws_answer[f'{col_letter}10'] = country
    ws_answer[f'{col_letter}11'] = value
    ws_answer[f'{col_letter}12'] = f"={col_letter}11*$C$2*$C$4/100"

# ----------------- STEP 3 -----------------
df_total = pd.read_excel(file_path, sheet_name='Total Reserves', header=7)
row_2025_idx = None
for idx, val in enumerate(df_total.iloc[:, 0]):
    if str(val) == '2025':
        row_2025_idx = idx
        break

total_reserves_2025 = {}
if row_2025_idx is not None:
    for col in df_total.columns[2:]:
        country_name = col.split(':')[0].strip()
        val_2025 = df_total.at[row_2025_idx, col]
        if pd.notna(val_2025) and str(val_2025).strip() != '':
            total_reserves_2025[country_name] = float(val_2025)

curr_col = 3 # Column C

for i, (country, value) in enumerate(countries_with_value_2025.items()):
    if country in total_reserves_2025:
        col_letter = get_col_letter(curr_col)
        # Previous Step 2 col_letter for this country:
        step2_col_letter = get_col_letter(start_col + i)
        
        ws_answer[f'{col_letter}19'] = country
        ws_answer[f'{col_letter}20'] = f"={step2_col_letter}11"
        ws_answer[f'{col_letter}21'] = f"={step2_col_letter}12"
        
        formula = f'=XLOOKUP({col_letter}19&"*", \'Total Reserves\'!$C$8:$R$8, \'Total Reserves\'!$C$17:$R$17, "", 2)'
        ws_answer[f'{col_letter}22'] = formula
        
        ws_answer[f'{col_letter}23'] = f"={col_letter}21/{col_letter}22*100"
        
        curr_col += 1

wb.save('/root/output/rar_result.xlsx')
print("All steps done")
