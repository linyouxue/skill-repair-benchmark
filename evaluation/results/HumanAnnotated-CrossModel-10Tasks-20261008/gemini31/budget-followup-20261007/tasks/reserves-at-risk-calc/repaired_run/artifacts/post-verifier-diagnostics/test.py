from openpyxl import load_workbook
wb = load_workbook('/root/output/rar_result.xlsx')
ws = wb['Answer']
print('Total Reserve formula check')
for col in range(3, 10):
    c = chr(64+col)
    print(ws.cell(row=23, column=col).value)
