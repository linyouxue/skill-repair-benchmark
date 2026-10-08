from openpyxl import load_workbook
file_path = '/root/output/rar_result.xlsx'
wb = load_workbook(file_path)
ws = wb['Answer']
ws['C3'] = "=NORM.S.INV(0.95)"
ws['C4'] = "='Gold price'!D430"
ws['C5'] = "=C4*SQRT(12/3)"
ws['C6'] = "='Gold price'!E430"
ws['C2'] = None

# Update formula in row 13 again:
for col in range(3, 20):
    c = chr(64+col)
    val = ws.cell(row=13, column=col).value
    if val:
        ws.cell(row=13, column=col).value = f"={c}12*$C$3*$C$5/100"

wb.save(file_path)
