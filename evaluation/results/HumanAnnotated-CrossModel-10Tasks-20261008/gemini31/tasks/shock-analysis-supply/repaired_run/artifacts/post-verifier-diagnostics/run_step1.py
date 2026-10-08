import pandas as pd
from openpyxl import load_workbook
import urllib.request

# Load test-supply.xlsx
wb = load_workbook('test-supply.xlsx')
ws_pwt = wb['PWT']

# Get PWT data
pwt_df = pd.read_excel('pwt1001.xlsx', sheet_name='Data')
geo_df = pwt_df[pwt_df['countrycode'] == 'GEO'].copy()

# The years we need in PWT sheet: Check the sheet
years_in_pwt = []
row_idx = 4
while ws_pwt.cell(row=row_idx, column=1).value is not None:
    years_in_pwt.append(ws_pwt.cell(row=row_idx, column=1).value)
    row_idx += 1

for i, year in enumerate(years_in_pwt):
    row_in_excel = i + 4
    if pd.isna(year) or str(year).strip() == '':
        break
    
    # Try to find the year in geo_df
    year_val = int(year)
    matched_row = geo_df[geo_df['year'] == year_val]
    if not matched_row.empty:
        rnna = matched_row['rnna'].values[0]
        labsh = matched_row['labsh'].values[0]
        ws_pwt.cell(row=row_in_excel, column=2).value = rnna
        ws_pwt.cell(row=row_in_excel, column=3).value = labsh

wb.save('test-supply.xlsx')
print("PWT data filled.")
