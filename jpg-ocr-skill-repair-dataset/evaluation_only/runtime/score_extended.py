"""Isolated grader: official batch score plus aggregate diagnostics, no gold leak."""
from __future__ import annotations

import argparse
import ast
import contextlib
import importlib.util
import io
import json
import re
from pathlib import Path

from prepare_batches import read_batches

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent


def load_checker(split, actual=None, batch=None):
    folder = ROOT / f"grader/{split}"
    if batch:
        if batch not in {row["batch_id"] for row in read_batches(split)}:
            raise ValueError("Unknown batch ID.")
        folder = folder / "batches" / batch
    path = folder / "test_outputs.py"
    spec = importlib.util.spec_from_file_location("official_jpg_checker", path)
    module = importlib.util.module_from_spec(spec)
    if actual is None:
        spec.loader.exec_module(module)
    else:
        class RedirectLocalOutput(ast.NodeTransformer):
            def visit_Constant(self, node):
                if node.value == "/app/workspace/stat_ocr.xlsx":
                    return ast.copy_location(ast.Constant(str(actual)), node)
                return node
        tree = RedirectLocalOutput().visit(ast.parse(path.read_text(encoding="utf-8")))
        ast.fix_missing_locations(tree)
        exec(compile(tree, str(path), "exec"), module.__dict__)
    return module


def cell_string(value):
    return "" if value is None else str(value)


STRATA = ["no_hint_conflict_detected", "hint_conflict_unique_payment", "literal_rule_ambiguous"]
HEADER = ["filename", "date", "total_amount"]


def read_predictions(actual, expected_names):
    predictions, issues, rows = {}, [], []
    if not actual.is_file():
        return predictions, ["output_missing"], rows
    workbook = None
    try:
        workbook = load_workbook(actual, data_only=True)
        if workbook.sheetnames != ["results"]:
            issues.append("invalid_sheets")
        else:
            rows = list(workbook["results"].values)
            if not rows or list(rows[0]) != HEADER:
                issues.append("invalid_header_or_columns")
            else:
                names = [cell_string(row[0]) for row in rows[1:]]
                if len(names) != len(expected_names):
                    issues.append("wrong_row_count")
                if len(names) != len(set(names)):
                    issues.append("duplicate_filenames")
                if set(names) != set(expected_names):
                    issues.append("missing_or_unexpected_filenames")
                if names != sorted(names):
                    issues.append("wrong_filename_order")
                for row in rows[1:]:
                    filename, date, amount = row
                    if not isinstance(filename, str) or not filename:
                        issues.append("invalid_filename_type")
                    # Null is a permitted extraction fallback, scored against GT.
                    if date is not None and (not isinstance(date, str) or not re.fullmatch(r"\d{4}-\d{2}-\d{2}", date)):
                        issues.append("invalid_date_format")
                    if amount is not None and (not isinstance(amount, str) or not re.fullmatch(r"\d+\.\d{2}", amount)):
                        issues.append("invalid_amount_format_or_type")
                    predictions[cell_string(filename)] = [cell_string(date), cell_string(amount)]
    except Exception:
        issues = ["unreadable_workbook"]
        predictions, rows = {}, []
    finally:
        if workbook is not None:
            workbook.close()
    return predictions, sorted(set(issues)), rows


def summarize(group):
    count = len(group)
    dates = sum(row["date_correct"] for row in group)
    amounts = sum(row["amount_correct"] for row in group)
    joint = sum(row["date_correct"] and row["amount_correct"] for row in group)
    return {"receipts": count, "fields": 2 * count, "date_correct": dates, "amount_correct": amounts,
            "field_correct": dates + amounts, "receipt_joint_correct": joint,
            "field_accuracy": (dates + amounts) / (2 * count) if count else None,
            "date_accuracy": dates / count if count else None,
            "amount_accuracy": amounts / count if count else None,
            "receipt_joint_accuracy": joint / count if count else None}


def score(actual, split, batch=None):
    checker = load_checker(split, actual, batch)
    expected = Path(checker.__file__).parent / "stat_oracle.xlsx"
    if not expected.is_file():
        raise RuntimeError("Missing expanded grader oracle; do not run an experiment.")
    manifest = json.loads((ROOT / "curator/frozen_manifest.json").read_text(encoding="utf-8"))
    expected_rows = [row for row in manifest if row["split"] == split]
    if batch:
        names = next(row["filenames"] for row in read_batches(split) if row["batch_id"] == batch)
        expected_rows = [row for row in expected_rows if row["filename"] in names]
    predictions, issues, _ = read_predictions(actual, [row["filename"] for row in expected_rows])
    reward = 0
    if actual.exists():
        if "unreadable_workbook" not in issues:
            try:
                # Complete unchanged official test, with stdout/expected text private.
                with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
                    checker.test_outputs()
                reward = 1
            except (AssertionError, FileNotFoundError):
                reward = 0
    observations = []
    for row in expected_rows:
        predicted = predictions.get(row["filename"], ["", ""])
        observations.append({"date_correct": predicted[0] == row["ground_truth"]["date"], "amount_correct": predicted[1] == row["ground_truth"]["total_amount"], "stratum": row["evaluation_stratum"]})

    summary = summarize(observations)
    return {"split": split, "batch_id": batch, "original_batch_reward": reward,
            "delivery_status": "valid" if not issues else "invalid", "delivery_issues": issues,
            "candidate_selection_eligible": not issues,
            "objective_field_accuracy": summary["field_accuracy"] if not issues else None,
            "all_receipts": summary,
            "strata": {stratum: summarize([row for row in observations if row["stratum"] == stratum]) for stratum in STRATA},
            "exposes_expected_values": False}


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--actual", required=True, type=Path)
    parser.add_argument("--split", required=True, choices=["train", "validation", "test"])
    parser.add_argument("--batch", help="Frozen batch ID; omit only for a full-split workbook.")
    parser.add_argument("--report", type=Path)
    args = parser.parse_args()
    result = score(args.actual, args.split, args.batch)
    text = json.dumps(result, ensure_ascii=False, indent=2) + "\n"
    if args.report:
        args.report.write_text(text, encoding="utf-8")
    else:
        print(text)
